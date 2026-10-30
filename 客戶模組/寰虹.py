# -*- coding: utf-8 -*-
"""
客戶模組：寰虹（掃描 PDF，要先辨識）
====================================
客戶給的是掃描的 PDF，一份 PDF 有好幾種表，每頁最上面是標題，例如：
  ・114年度薪資明細彙總表      → 每人：薪資／扣繳金額／伙食費；欄：113年十二月、一月～十一月、合計
  ・114年度各項獎金給付總表    → 欄：113年年終、端午、中秋、業績獎金、週年禮金、合計
  ・114年度免稅薪資（特休未休返薪）→ 免稅，不放
  ・114年度員工自提退休金明細  → 每人：投保級距／自提率／自提金額
  ・113年度員工紅利申請表      → 欄：一月～十二月；放進那個月的薪資總額
  ・尾牙抽獎及加碼紅包          → 91-1（競技競賽及機會中獎），另外出一本清冊，月份在畫面上選（預設 2 月）
放進清冊（以清冊 114 年為例）：
  ・薪資那列 → 那個月的薪資總額；113年十二月 → 前期薪資
  ・扣繳金額 → 扣繳稅額
  ・伙食費   → 每人多一列「伙食費」，放在同一個月
  ・獎金：上一年度的年終 → 上期年終；端午 → 端午獎金；中秋 → 中秋獎金；其他（業績、週年…）→ 端午獎金
辨識用 RapidOCR，在這台電腦上離線跑（不會上網）。需要：pip install rapidocr_onnxruntime pypdfium2 opencc-python-reimplemented
"""
import difflib
import os
import re
import 清冊共用 as 共

名稱 = "寰虹（掃描 PDF）"
說明 = "掃描的 PDF，每頁一種表（薪資、獎金、自提、紅利、尾牙），先辨識再整理"
多個薪資檔 = True
預設客戶 = ["寰虹實業"]
預設設定 = {
    "年度": 114,
    "薪資總額": "薪資", "退休自提": "自提金額", "扣繳稅額": "扣繳金額",
    "掃描檔們": [],
    "尾牙月份": 2,
    "執行業務另開": False,
}

# ---- 認表的關鍵字（標題裡有這些字 → 這種表）。辨識常把「獎」認成「樊」，所以兩個都放
表種類 = [
    ("免稅", ["免稅"]),
    ("自提", ["自提", "退休"]),
    ("91-1", ["尾牙", "抽獎", "抽樊", "加碼紅包"]),
    ("紅利", ["紅利"]),
    ("獎金", ["各項", "給付總表"]),
    ("薪資", ["薪資明細", "薪資", "明細"]),
]
# ---- 每人那幾列（左邊的字）→ 用途
列名認法 = [("伙食費", ["伙食", "伙會", "伙食費"]), ("稅", ["扣繳", "扣緻", "扣繳金額"]),
         ("自提", ["自提金額", "自提金"]), ("級距", ["投保", "級距"]), ("自提率", ["自提率", "提率"]),
         ("蓋章", ["蓋章", "盞章", "藍章", "蓋", "盞"]), ("身分", ["身分", "身份"]),
         ("總額", ["薪資", "薪", "獎金", "樊金", "免稅薪資", "金"])]
# ---- 獎金表的欄（表頭的字）→ 清冊的欄
獎金欄 = [("年終", "上期年終"), ("端午", "端午獎金"), ("中秋", "中秋獎金")]   # 沒對到的 → 端午獎金
其他獎金放 = "端午獎金"
月名 = ["一月", "二月", "三月", "四月", "五月", "六月", "七月", "八月", "九月", "十月", "十一月", "十二月"]


# ================================================================ 辨識（PDF／圖片 → 文字框）
_快取 = {}
_引擎 = {}


def 頁數(路徑):
    if os.path.splitext(路徑)[1].lower() != ".pdf":
        return 1
    try:
        import pypdfium2 as pdfium
        return len(pdfium.PdfDocument(路徑))
    except Exception:
        return 1


def 辨識檔案(路徑, 進度=None):
    """回傳每頁 {"寬","高","框":[{字,x1,x2,y1,y2,分}]}；同一個檔沒改過就不重新辨識"""
    鍵 = (路徑, os.path.getmtime(路徑))
    if 鍵 in _快取:
        return _快取[鍵]
    try:
        import numpy as np
        from rapidocr_onnxruntime import RapidOCR
    except ImportError:
        raise SystemExit("找不到辨識套件，請先執行：pip install rapidocr_onnxruntime pypdfium2 opencc-python-reimplemented")
    if "ocr" not in _引擎:
        _引擎["ocr"] = RapidOCR()
        try:
            import opencc
            _引擎["繁"] = opencc.OpenCC("s2tw").convert
        except ImportError:
            _引擎["繁"] = lambda s: s
    頁們 = []
    for 第幾, img in enumerate(_讀頁面(路徑)):
        if 進度:
            進度(第幾)                              # 告訴畫面：開始辨識第幾頁（從 0 開始）
        結果, _ = _引擎["ocr"](img)
        框們 = []
        for 點, 字, 分 in 結果 or []:
            xs = [p[0] for p in 點]; ys = [p[1] for p in 點]
            框們.append({"字": _引擎["繁"](str(字)).strip(), "x1": min(xs), "x2": max(xs),
                       "y1": min(ys), "y2": max(ys), "分": float(分)})
        頁 = {"寬": img.shape[1], "高": img.shape[0], "框": 框們}
        r = 解析頁(頁)
        for x0, y0, x1, y1 in (r or {}).get("缺名", []):    # 人名沒辨識到：把那一小塊放大 2 倍再辨識一次
            塊 = img[max(0, y0):y1, x0:x1]
            if 塊.size == 0:
                continue
            import cv2
            塊 = cv2.copyMakeBorder(cv2.resize(塊, None, fx=2, fy=2, interpolation=cv2.INTER_CUBIC),
                                   20, 20, 20, 20, cv2.BORDER_CONSTANT, value=(255, 255, 255))
            結果2, _ = _引擎["ocr"](塊)
            字們 = [_引擎["繁"](str(t)).strip() for _, t, _ in 結果2 or []]
            字們 = [共.清字(t) for t in 字們 if re.fullmatch(r"[\u4e00-\u9fff]{2,5}", 共.清字(t))]
            if 字們:
                框們.append({"字": 字們[0], "x1": x0 + 10, "x2": x1 - 10, "y1": y0 + 15, "y2": y1 - 15, "分": 0.5, "補": True})
        頁["補過"] = True
        r = 解析頁(頁)
        for 人 in (r or {}).get("人們", []):           # 人名常常少一個字：放大 2 倍、3 倍各辨識一次，三個結果投票
            b = 人["框"]
            x0, x1, y0, y1 = int(b["x1"]) - 15, int(b["x2"]) + 15, int(b["y1"]) - 10, int(b["y2"]) + 10
            塊 = img[max(0, y0):y1, max(0, x0):x1]
            if 塊.size == 0:
                continue
            import cv2
            票 = [共.清字(b["字"])]
            for 倍 in (2, 3):
                放大 = cv2.copyMakeBorder(cv2.resize(塊, None, fx=倍, fy=倍, interpolation=cv2.INTER_CUBIC),
                                       20, 20, 20, 20, cv2.BORDER_CONSTANT, value=(255, 255, 255))
                結果3, _ = _引擎["ocr"](放大)
                for _, t, _ in 結果3 or []:
                    t = 共.清字(_引擎["繁"](str(t)))
                    if re.fullmatch(r"[\u4e00-\u9fff]{2,5}", t):
                        票.append(t)
                        break
            from collections import Counter
            數 = Counter(票)
            b["字"] = max(數, key=lambda t: (數[t], len(t)))     # 票最多的；一樣多就選字多的（少字比多字常見）
        頁們.append(頁)
    _快取[鍵] = 頁們
    return 頁們


def _讀頁面(路徑):
    import numpy as np
    副 = os.path.splitext(路徑)[1].lower()
    if 副 == ".pdf":
        try:
            import pypdfium2 as pdfium
        except ImportError:
            raise SystemExit("找不到 PDF 套件，請先執行：pip install pypdfium2")
        pdf = pdfium.PdfDocument(路徑)
        for i in range(len(pdf)):
            yield np.array(pdf[i].render(scale=200 / 72).to_pil().convert("RGB"))[:, :, ::-1].copy()
    else:
        import cv2
        img = cv2.imdecode(np.fromfile(路徑, dtype=np.uint8), cv2.IMREAD_COLOR)
        if img is None:
            raise SystemExit(f"圖片讀不到：{os.path.basename(路徑)}")
        yield img


# ================================================================ 解析一頁
數字樣式 = re.compile(r"^-?[\d,，.．]+%?$")


def 轉數字(字):
    字 = 字.strip().rstrip(".．,，")
    if re.fullmatch(r"[-－—_]+", 字):
        return 0
    if not 數字樣式.match(字) or "%" in 字:
        return None
    數 = re.sub(r"[^\d]", "", 字)
    return int(數) if 數 else None


def 中心(b):
    return (b["x1"] + b["x2"]) / 2, (b["y1"] + b["y2"]) / 2


def 表頭記號(字):
    """表頭一個框可能黏了好幾欄（例如「113年十二月一月」）→ 拆成 [記號]"""
    記號 = []
    for m in re.finditer(r"(\d{2,3})年(十二|十一|十|[一二三四五六七八九])月|(十二|十一|十|[一二三四五六七八九])月|合計|"
                         r"(\d{2,3})年年終|端午|中秋|業績|[週調]年|禮金|年終|抽[獎樊]|加碼|紅包", 字):
        記號.append(m.group(0))
    return 記號


def 解析頁(頁):
    框們 = 頁["框"]
    if not 框們:
        return None
    # ---- 標題：最上面的字
    最上 = min(b["y1"] for b in 框們)
    標題框 = [b for b in 框們 if b["y1"] < 最上 + 40]
    標題 = "".join(b["字"] for b in sorted(標題框, key=lambda b: b["x1"]))
    類 = next((k for k, 詞們 in 表種類 if any(w in 標題 for w in 詞們)), None)
    年 = re.search(r"(\d{3})年", 標題)
    # ---- 表頭：有「姓名」「合計」的那一列
    合計框 = next((b for b in 框們 if b["字"] == "合計" or b["字"].endswith("合計")), None)
    姓名框 = next((b for b in 框們 if "姓名" in b["字"]), None)
    if 姓名框 is None and 合計框 is not None:          # 「姓名」沒辨識到：用「NO」或左邊的列名推一個位置
        _, y合0 = 中心(合計框)
        同列 = [b for b in 框們 if abs(中心(b)[1] - y合0) < 30]
        NO = next((b for b in 同列 if 共.清字(b["字"]).upper() in ("NO", "N0", "NO.")), None)
        列名們 = [b for b in 框們 if 認列名(b["字"]) in ("總額", "稅", "伙食費", "級距", "自提率", "自提") and 中心(b)[1] > y合0]
        左列名 = min((b["x1"] for b in 列名們), default=None)
        if NO is not None or 左列名 is not None:
            x0 = NO["x2"] + 40 if NO is not None else 左列名 - 220
            x1 = 左列名 - 40 if 左列名 is not None else x0 + 120
            y = 中心(NO)[1] if NO is not None else y合0
            姓名框 = {"字": "姓名", "x1": x0, "x2": max(x0 + 20, x1), "y1": y - 10, "y2": y + 10}
    if 姓名框 is None or 合計框 is None:
        return {"標題": 標題, "類": 類, "年": int(年.group(1)) if 年 else None, "人們": [], "錯": "找不到表頭（姓名、合計）"}
    _, 表頭y = 中心(姓名框)
    (x合, y合) = 中心(合計框)
    (x姓, y姓) = 中心(姓名框)
    斜 = (y合 - y姓) / max(1, x合 - x姓)                # 掃描有一點歪：每往右 1 點，往下幾點
    正y = lambda b: 中心(b)[1] - 斜 * 中心(b)[0]
    表頭框 = [b for b in 框們 if abs(正y(b) - 正y(姓名框)) < 25 and b["x1"] > 姓名框["x2"]]
    # 表頭記號 → 欄（名稱, x 中心）
    欄們 = []
    for b in sorted(表頭框, key=lambda b: b["x1"]):
        記 = 表頭記號(b["字"])
        if not 記:                               # 認不出的表頭（例如「抽獎獎金」辨識成「抽金」）→ 整格當一欄
            字 = 共.清字(b["字"])
            if not re.search(r"[\u4e00-\u9fff]", 字) or 字 in ("姓名", "NO"):
                continue
            記 = [字]
        寬 = (b["x2"] - b["x1"]) / len(記)
        for i, k in enumerate(記):
            欄們.append([k, b["x1"] + 寬 * (i + 0.5)])
    欄們 = 補月份(欄們)
    # ---- 左邊的列名（薪資／扣繳金額／伙食費…）和姓名
    數字框 = [b for b in 框們 if 轉數字(b["字"]) is not None and 正y(b) > 正y(姓名框) + 15 and b["x1"] > 姓名框["x2"]
           and (轉數字(b["字"]) >= 10 or 轉數字(b["字"]) == 0)]       # 只有一位數的多半是雜點，不算
    第一欄x = min((x for _, x in 欄們), default=x合)
    列名框 = []
    for b in 框們:
        x, _ = 中心(b)
        if 正y(b) <= 正y(姓名框) + 15 or x >= 第一欄x - 30 or x <= x姓 + 20:
            continue
        用 = 認列名(b["字"])
        if 用:
            列名框.append((用, 正y(b), b))
    列名x = sorted(中心(b)[0] for _, _, b in 列名框)[len(列名框) // 2] if 列名框 else x姓 + 150
    名框們 = []
    for b in 框們:
        x, _ = 中心(b)
        字 = 共.清字(b["字"])
        if 正y(b) <= 正y(姓名框) + 15 or abs(x - x姓) > 列名x - x姓 - 20 or not re.fullmatch(r"[一-鿿]{2,5}", 字):
            continue
        if 認列名(字) or any(w in 字 for w in ("秘", "祕", "密", "住址", "編號", "身分", "身份")):
            continue
        y = 正y(b)
        # 姓名旁邊（同一列）要有「第一列」的列名或數字，才是人名（地址被切斷的字不會有）
        if any(abs(yy - y) < 18 for 用, yy, _ in 列名框 if 用 == "總額") or \
                any(abs(正y(n) - y) < 18 for n in 數字框):
            名框們.append((字, y, b))
    名框們.sort(key=lambda t: t[1])
    缺名 = []                                   # 有第一列（列名）但旁邊沒辨識到人名 → 這個範圍要放大再辨識一次
    第一列名 = {"總額", "級距"}
    for 用, yy, b in 列名框:
        if 用 in 第一列名 and not any(abs(y - yy) < 18 for _, y, _ in 名框們):
            x0, x1 = 姓名框["x1"] - 60, b["x1"] - 10
            yc = 中心(b)[1]
            缺名.append((int(max(0, x0)), int(yc - 28), int(x1), int(yc + 22)))
    if 缺名 and 頁.get("補過") is None:
        return {"標題": 標題, "類": 類, "缺名": 缺名}
    if not 名框們:
        return {"標題": 標題, "類": 類, "年": int(年.group(1)) if 年 else None, "人們": [], "欄們": 欄們, "錯": "找不到人名"}
    # ---- 每一組（一個人）各列的位置：用找到的列名算出「離第一列多遠」（這頁每組都一樣）
    距離 = {}
    for i, (名, y0, _) in enumerate(名框們):
        y下 = 名框們[i + 1][1] if i + 1 < len(名框們) else 1e9
        for 用, yy, _ in 列名框:
            if y0 - 18 <= yy < y下 - 18 and 用 not in ("總額", "身分"):
                距離.setdefault(用, []).append(yy - y0)
    距離 = {k: sorted(v)[len(v) // 2] for k, v in 距離.items()}
    人們 = []
    for i, (名, y0, _) in enumerate(名框們):
        y下 = 名框們[i + 1][1] if i + 1 < len(名框們) else 1e9
        列們 = {k: y0 + d for k, d in 距離.items()}
        if not any(abs(d) < 12 for d in 距離.values()):    # 第一列沒有別的列名（例如投保級距）→ 第一列就是金額
            列們["總額"] = y0
        值 = {}
        for b in 數字框:
            y = 正y(b)
            if not (y0 - 18 <= y < y下 - 18):
                continue
            用 = min(列們, key=lambda k: abs(列們[k] - y))
            if abs(列們[用] - y) > 20:
                continue
            欄 = 對欄(b, 欄們)
            if 欄 is None:
                continue
            值.setdefault(用, {})[欄] = 值.setdefault(用, {}).get(欄, 0) + 轉數字(b["字"])
        人們.append({"名": 名, "值": 值, "框": 名框們[i][2]})
    錯 = "有人名沒辨識到，請對照 PDF" if 缺名 else ""
    return {"標題": 標題, "類": 類, "年": int(年.group(1)) if 年 else None, "欄們": 欄們, "人們": 人們, "錯": 錯}


def 認列名(字):
    字 = 共.清字(字)
    if not 字 or len(字) > 6:
        return None
    for 用, 詞們 in 列名認法:
        if any(w == 字 or (len(w) >= 2 and w in 字) for w in 詞們):
            return 用
    return None


# 月份欄有的沒辨識到、有的黏在一起認錯（例如「十月十月」）→ 照順序重新排：113年十二月、一月、二月…十二月、合計
def 補月份(欄們):
    月 = [(i, k, x) for i, (k, x) in enumerate(欄們) if re.search(r"月$", k)]
    if len(月) < 3:
        return 欄們
    def 號(k):
        m = re.search(r"(十二|十一|十|[一二三四五六七八九])月$", k)
        return 月名.index(m.group(1) + "月") + 1
    有去年 = any(re.match(r"\d{2,3}年", k) for _, k, _ in 月)
    序 = (["前12"] if 有去年 else []) + [f"{m}" for m in range(1, 13)]
    # 用辨識到的月份（不重複的）算出「第幾欄 → x」的直線
    點 = []
    看過 = set()
    for _, k, x in 月:
        n = "前12" if re.match(r"\d{2,3}年", k) else str(號(k))
        if n in 看過:
            continue
        看過.add(n)
        點.append((序.index(n), x))
    if len(點) < 2:
        return 欄們
    n = len(點)
    mx = sum(a for a, _ in 點) / n; my = sum(b for _, b in 點) / n
    斜 = sum((a - mx) * (b - my) for a, b in 點) / max(1e-9, sum((a - mx) ** 2 for a, _ in 點))
    截 = my - 斜 * mx
    結果 = [[("去年十二月" if s == "前12" else 月名[int(s) - 1]), 截 + 斜 * i] for i, s in enumerate(序)]
    合計 = [[k, x] for k, x in 欄們 if k == "合計"]
    結果 = [r for r in 結果 if not 合計 or r[1] < 合計[0][1] - 斜 * 0.5]
    return 結果 + 合計 + [[k, x] for k, x in 欄們 if not re.search(r"月$", k) and k != "合計"]


def 對欄(b, 欄們):
    """數字是靠右對齊的 → 右邊界落在「這欄表頭中心 ～ 下一欄表頭中心」之間，就是這一欄"""
    排 = sorted(欄們, key=lambda t: t[1])
    if not 排:
        return None
    右 = b["x2"]
    前 = [k for k, x in 排 if x < 右]
    return 前[-1] if 前 else None


# ================================================================ 整份 PDF → 每頁的表（人名統一、驗算）
def 讀掃描檔(路徑們, 名冊名們=(), 進度=None):
    """回傳 (頁們, 警告)；頁們：[{"檔","頁","類","標題","年","欄們","人們"}]"""
    頁們, 警告 = [], []
    for 路徑 in 路徑們:
        try:
            原頁們 = 辨識檔案(路徑, (lambda i, p=路徑: 進度(p, i)) if 進度 else None)
        except SystemExit:
            raise
        except Exception as e:
            警告.append(f"{os.path.basename(路徑)} 辨識失敗：{e}")
            continue
        前類 = None
        for i, 頁 in enumerate(原頁們):
            r = 解析頁(頁)
            if not r:
                continue
            r.update(檔=os.path.basename(路徑), 頁=i + 1)
            if r.get("類") is None:                # 沒有標題的頁 → 跟前一頁同一種表（表格接續到下一頁）
                r["類"] = 前類
            前類 = r["類"]
            if r.get("錯"):
                警告.append(f"{r['檔']} 第 {i + 1} 頁：{r['錯']}")
            頁們.append(r)
    統一人名(頁們, 名冊名們, 警告)
    for r in 頁們:                                 # 驗算：每一列的各欄加起來 ＝ 合計
        for 人 in r.get("人們", []):
            for 用, 值 in 人["值"].items():
                if "合計" in 值 and 用 != "自提率":
                    其他 = sum(v for k, v in 值.items() if k != "合計")
                    if 其他 != 值["合計"]:
                        警告.append(f"{r['檔']} 第 {r['頁']} 頁 {人['名']}「{列中文(用)}」：各欄加起來 {其他:,} ≠ 合計 {值['合計']:,}"
                                  "（可能有數字辨識錯，請對照 PDF）")
    return 頁們, 警告


def 列中文(用):
    return {"總額": "金額", "稅": "扣繳金額", "伙食費": "伙食費", "自提": "自提金額", "級距": "投保級距"}.get(用, 用)


def 統一人名(頁們, 名冊名們, 警告):
    """辨識錯的人名（例如「張張介禎」「昭蓉」）→ 換成出現比較多次的寫法，或人員名冊上的名字"""
    from collections import Counter
    次數 = Counter(人["名"] for r in 頁們 for 人 in r.get("人們", []))
    標準 = set(名冊名們)
    for 名, n in 次數.most_common():
        if 名 in 標準:
            continue
        像 = difflib.get_close_matches(名, list(標準), n=1, cutoff=0.6)
        if 像 and (像[0] in 名冊名們 or 次數[像[0]] >= n):
            continue
        標準.add(名)
    換 = {}
    for 名 in 次數:                              # 少字：「張介」是「張介禎」的一部分 → 當成張介禎
        長 = [x for x in 次數 if x != 名 and len(x) == len(名) + 1 and all(ch in x for ch in 名) and x.startswith(名[0])]
        if len(長) == 1 and 名 not in 名冊名們:
            換[名] = 長[0]
    for 名 in 次數:
        if 名 in 換:
            continue
        if 名 in 名冊名們:
            continue
        候選 = [x for x in 標準 if x != 名]
        像 = difflib.get_close_matches(名, 候選, n=1, cutoff=0.5)
        if 像 and (像[0] in 名冊名們 or 次數[像[0]] > 次數[名] or (名 not in 標準)):
            換[名] = 像[0]
    for r in 頁們:
        for 人 in r.get("人們", []):
            if 人["名"] in 換:
                警告.append(f"{r['檔']} 第 {r['頁']} 頁：人名辨識成「{人['名']}」，當成「{換[人['名']]}」，請確認")
                人["名"] = 換[人["名"]]


def 月份欄(k, 年):
    """表頭的月份 → 清冊的欄：去年十二月 → 前期薪資；一月～十二月 → 一月～十二月"""
    if k == "去年十二月":
        return "前期薪資"
    return k if k in 共.中文月 else None


def 獎金清冊欄(k, 年):
    m = re.match(r"(\d{2,3})年", k)
    for 詞, 欄 in 獎金欄:
        if 詞 in k:
            if 詞 == "年終" and m and int(m.group(1)) != 年 - 1:
                return 其他獎金放
            return 欄
    return 其他獎金放


# ---------------------------------------------------------------- 讀取（產生清冊時共用核心會叫這個）
def 讀取(設, 年, wb, wbf, 清冊欄):
    檔們 = [p for p in 設.get("掃描檔們", []) if p]
    if not 檔們:
        raise SystemExit("請在「檔案與年度」頁加入客戶的掃描 PDF")
    缺表 = [os.path.basename(p) for p in 檔們 if not os.path.exists(p)]
    名冊名們 = []
    if 設.get("人員名冊") and os.path.exists(設["人員名冊"]):
        名冊名們 = list(共.讀前年度(設["人員名冊"]))
    頁們, 警告 = 讀掃描檔([p for p in 檔們 if os.path.exists(p)], 名冊名們)
    來源們 = {}

    def 來源(鍵, 欄位, 類型, 顯示, **屬性):
        if 鍵 not in 來源們:
            t = 共.資料來源(欄位, 欄位, 類型)               # 名稱用欄位（摘要寫「月薪表：一月～十一月」）
            t.顯示名稱 = 顯示
            t.總額, t.自提, t.扣繳 = [], None, []
            for k, v in 屬性.items():
                setattr(t, k, v)
            來源們[鍵] = t
        return 來源們[鍵]

    免稅 = []
    尾牙月 = int(設.get("尾牙月份") or 2)
    for r in 頁們:
        頁名 = f"{r['檔']} 第{r['頁']}頁"
        類 = r.get("類")
        if 類 is None:
            警告.append(f"{頁名}：認不出是哪一種表（標題：{r.get('標題', '')}），沒有讀")
            continue
        if 類 == "免稅":
            免稅.append(f"{頁名}（{len(r['人們'])} 人，{sum(人['值'].get('總額', {}).get('合計', 0) for 人 in r['人們']):,}）")
            continue
        for 人 in r["人們"]:
            名, 值 = 人["名"], 人["值"]
            if 類 in ("薪資", "紅利", "自提"):
                for k in r["欄們"]:
                    欄 = 月份欄(k[0], 年)
                    if not 欄:
                        continue
                    if 欄 not in 清冊欄:
                        raise SystemExit(f"清冊樣板沒有「{欄}」欄")
                    if 類 == "自提":
                        金 = 值.get("自提", {}).get(k[0], 0)
                        if 金:
                            t = 來源(("自提", 欄), 欄, "月薪", f"{頁名} 自提退休金明細（{欄}）", 併入=True)
                            t.加(名, 自提=金)
                        continue
                    總 = 值.get("總額", {}).get(k[0], 0)
                    稅 = 值.get("稅", {}).get(k[0], 0)
                    if 總 or 稅:
                        標 = "員工紅利" if 類 == "紅利" else "薪資"
                        t = 來源((標, 欄), 欄, "月薪", f"{頁名} {標}（{欄}）", 併入=(類 == "紅利"))
                        t.加(名, 總額=總, 稅=稅)
                    伙 = 值.get("伙食費", {}).get(k[0], 0)
                    if 伙:
                        t = 來源(("伙食費", 欄), 欄, "月薪", f"{頁名} 伙食費（{欄}）", 明細列="伙食費", 併入=True)
                        t.加(名, 總額=伙)
            elif 類 == "獎金":
                for k in r["欄們"]:
                    if k[0] == "合計":
                        continue
                    欄 = 獎金清冊欄(k[0], 年)
                    if 欄 not in 清冊欄:
                        raise SystemExit(f"清冊樣板沒有「{欄}」欄")
                    總 = 值.get("總額", {}).get(k[0], 0)
                    稅 = 值.get("稅", {}).get(k[0], 0)
                    if 總 or 稅:
                        t = 來源(("獎金", 欄), 欄, "獎金", f"{頁名} 各項獎金（{欄}）")
                        t.加(名, 總額=總, 稅=稅)
            elif 類 == "91-1":
                總 = 值.get("總額", {})
                金 = 總.get("合計", sum(v for k, v in 總.items() if k != "合計"))
                稅 = sum(v for k, v in 值.get("稅", {}).items() if k != "合計")
                if 金 or 稅:
                    欄 = 共.中文月[尾牙月 - 1]
                    t = 來源(("91-1", 欄), 欄, "月薪", f"{頁名} 尾牙抽獎及加碼紅包（{尾牙月}月）",
                           清冊="91-1", 清冊標題="競技競賽及機會中獎（91-1）印領清冊", 併入=True)
                    t.加(名, 總額=金, 稅=稅)
    if 免稅:
        警告.append("免稅薪資（特休未休返薪）沒有放進清冊：" + "、".join(免稅))
    if not 來源們:
        raise SystemExit("掃描檔裡沒有讀到可以放進清冊的資料")
    return list(來源們.values()), 缺表, 警告


# ---------------------------------------------------------------- 把辨識結果存成 Excel（對照 PDF 檢查用）
def 匯出辨識結果(頁們, 路徑):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for r in 頁們:
        ws = wb.create_sheet(f"第{r['頁']}頁"[:31])
        ws.append([f"{r['檔']} 第 {r['頁']} 頁：{r.get('標題', '')}（{r.get('類') or '認不出'}）"])
        ws["A1"].font = Font(bold=True, size=12)
        欄名們 = [k for k, _ in r.get("欄們", [])]
        ws.append(["姓名", "列"] + 欄名們)
        for c in ws[2]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="C49A7C")
        for 人 in r.get("人們", []):
            for 用, 值 in 人["值"].items():
                ws.append([人["名"], 列中文(用)] + [值.get(k) for k in 欄名們])
        ws.column_dimensions["A"].width = 12
        ws.column_dimensions["B"].width = 10
    wb.save(路徑)


# ---------------------------------------------------------------- 畫面：檔案頁下面多「掃描的 PDF」和「91-1 月份」
def 頁面(介, 設):
    tk = 介.tk
    from tkinter import filedialog
    檔們 = list(設.get("掃描檔們", []))
    尾牙月 = tk.StringVar(value=f"{int(設.get('尾牙月份') or 2)}月")
    結果 = {"頁們": [], "警告": []}

    介.區塊標題(介.檔案頁, "客戶的掃描 PDF",
             右邊=lambda 框: 介.文字按鈕(框, "＋ 加入檔案…", lambda: 加入(), 色=介.PRIMARY_DARK, 移上色=介.TEXT,
                                     bg=介.BG).pack(side="right"))
    卡 = 介.卡片(介.檔案頁)
    清單 = tk.Frame(卡, bg=介.CARD_BG)
    清單.pack(fill="x")
    清單.columnconfigure(0, weight=1)
    下 = tk.Frame(卡, bg=介.CARD_BG)
    下.pack(fill="x", pady=(8, 0))
    介.標籤(下, "尾牙抽獎、加碼紅包（91-1）放在").pack(side="left")
    cb = 介.下拉(下, [f"{m}月" for m in range(1, 13)], 唯讀=True, 寬=5)
    cb.configure(textvariable=尾牙月)
    cb.pack(side="left", padx=(8, 0))
    介.文字按鈕(下, "匯出辨識結果…", lambda: 匯出(), 色=介.PRIMARY_DARK).pack(side="right")
    介.小字(卡, "加入檔案後會在這台電腦上辨識（離線，一頁約 3 秒）。每頁看標題認是哪一種表；人名辨識錯的會自動對成同一個人，"
              "各欄加起來跟合計對不起來的會標出來。「匯出辨識結果」可以存成 Excel 對照 PDF 檢查。", 寬=620).pack(fill="x", pady=(6, 0))
    狀 = {"文": "", "色": 介.TEXT_GRAY}
    種類名 = {"薪資": "薪資明細", "獎金": "各項獎金", "免稅": "免稅薪資（不放）", "自提": "自提退休金", "紅利": "員工紅利",
            "91-1": "尾牙抽獎（91-1，另一本清冊）", None: "認不出"}

    忙 = {"是": False}

    # 辨識：另外開一條執行緒在背後辨識，畫面不會當住；跳出一個小視窗顯示進度
    def 辨識():
        結果["頁們"], 結果["警告"] = [], []
        好的 = [p for p in 檔們 if os.path.exists(p)]
        if not 好的 or 忙["是"]:
            return
        import threading, queue
        名冊 = 介.收集().get("人員名冊")
        總頁 = sum(頁數(p) for p in 好的)
        前面 = {p: sum(頁數(q) for q in 好的[:k]) for k, p in enumerate(好的)}
        訊息 = queue.Queue()
        忙["是"] = True
        # ---- 進度視窗
        窗 = tk.Toplevel(介.root, bg=介.CARD_BG, padx=24, pady=18)
        窗.title("辨識中")
        窗.transient(介.root)
        窗.resizable(False, False)
        窗.protocol("WM_DELETE_WINDOW", lambda: None)        # 辨識中不能關
        tk.Label(窗, text="正在辨識掃描檔（在這台電腦上，離線）", font=介.字["標籤"], fg=介.TEXT, bg=介.CARD_BG).pack(anchor="w")
        文 = tk.Label(窗, text="準備中…", font=介.字["內文"], fg=介.TEXT_GRAY, bg=介.CARD_BG, anchor="w",
                     justify="left", wraplength=380)
        文.pack(anchor="w", pady=(8, 6))
        條 = 介.ttk.Progressbar(窗, length=380, mode="determinate", maximum=max(1, 總頁))
        條.pack(fill="x")
        秒字 = tk.Label(窗, text="", font=介.字["小"], fg=介.TEXT_GRAY, bg=介.CARD_BG, anchor="w")
        秒字.pack(anchor="w", pady=(6, 0))
        import time
        開始 = time.time()
        估 = {}
        窗.update_idletasks()                         # 視窗放在主視窗正中間
        x = 介.root.winfo_rootx() + (介.root.winfo_width() - 窗.winfo_reqwidth()) // 2
        y = 介.root.winfo_rooty() + (介.root.winfo_height() - 窗.winfo_reqheight()) // 3
        窗.geometry(f"+{max(0, x)}+{max(0, y)}")

        def 工作():                                  # 背景執行緒：只做辨識，不碰畫面
            try:
                名冊名們 = list(共.讀前年度(名冊)) if 名冊 and os.path.exists(名冊) else []
                r = 讀掃描檔(好的, 名冊名們, lambda p, i: 訊息.put(("頁", p, i)))
                訊息.put(("好", r))
            except SystemExit as ex:
                訊息.put(("錯", f"✗ {ex.code}"))
            except Exception as ex:
                訊息.put(("錯", f"✗ 辨識失敗：{ex}"))

        def 看進度():                                # 畫面這邊每 0.1 秒看一次進度
            try:
                while True:
                    種, *內容 = 訊息.get_nowait()
                    if 種 == "頁":
                        p, i = 內容
                        第 = 前面[p] + i + 1
                        條["value"] = 第 - 1
                        文.configure(text=f"第 {第}／{總頁} 頁\n{os.path.basename(p)}（這個檔的第 {i + 1} 頁）")
                        if 第 > 1:
                            估["每頁"] = (time.time() - 開始) / (第 - 1)
                            估["剩"] = 估["每頁"] * (總頁 - 第 + 1)
                            估["時"] = time.time()
                    else:
                        if 種 == "好":
                            結果["頁們"], 結果["警告"] = 內容[0]
                        else:
                            結果["警告"] = [內容[0]]
                        忙["是"] = False
                        窗.destroy()
                        try:
                            畫()
                        except tk.TclError:              # 辨識中換了客戶，這一頁已經不在了
                            pass
                        return
            except queue.Empty:
                pass
            if 窗.winfo_exists():
                if 估:
                    剩 = max(1, int(估["剩"] - (time.time() - 估["時"])))
                    秒字.configure(text=f"已經 {int(time.time() - 開始)} 秒，大約還要 {剩} 秒")
                else:
                    秒字.configure(text=f"已經 {int(time.time() - 開始)} 秒（第一頁會比較久）")
                介.root.after(100, 看進度)
        threading.Thread(target=工作, daemon=True).start()
        介.root.after(100, 看進度)

    def 畫():
        for w in 清單.winfo_children():
            w.destroy()
        r = 0
        for 路徑 in 檔們:
            tk.Label(清單, text=os.path.basename(路徑), font=介.字["標籤"], fg=介.TEXT, bg=介.CARD_BG,
                     anchor="w").grid(row=r, column=0, sticky="w", pady=(6, 2))
            介.文字按鈕(清單, "✕", lambda p=路徑: 移除(p), 移上色=介.ERR).grid(row=r, column=1, sticky="e")
            r += 1
            if not os.path.exists(路徑):
                tk.Label(清單, text="　✗ 找不到這個檔", font=介.字["小"], fg=介.ERR, bg=介.CARD_BG).grid(row=r, column=0, sticky="w")
                r += 1
                continue
            for 頁 in [x for x in 結果["頁們"] if x["檔"] == os.path.basename(路徑)]:
                文 = f"　第 {頁['頁']} 頁　{種類名.get(頁.get('類'), '認不出')}（{len(頁.get('人們', []))} 人）"
                tk.Label(清單, text=文, font=介.字["小"], fg=介.OK if 頁.get("類") else 介.WARN, bg=介.CARD_BG,
                         anchor="w").grid(row=r, column=0, sticky="w")
                r += 1
        要看 = [w for w in 結果["警告"] if not w.startswith("免稅")]
        if 要看:
            tk.Label(清單, text="⚠ " + "\n⚠ ".join(要看), font=介.字["小"], fg=介.WARN, bg=介.CARD_BG, anchor="w",
                     justify="left", wraplength=600).grid(row=r, column=0, columnspan=2, sticky="w", pady=(6, 0))
        if not 檔們:
            tk.Label(清單, text="還沒有加入 PDF", font=介.字["小"], fg=介.ERR, bg=介.CARD_BG).grid(row=0, column=0, sticky="w")
            狀.update(文="✗ 還沒加入 PDF", 色=介.ERR)
        elif any(w.startswith("✗") for w in 要看):
            狀.update(文="✗ 辨識有問題", 色=介.ERR)
        elif 要看:
            狀.update(文=f"⚠ {len(要看)} 個地方要確認", 色=介.WARN)
        else:
            狀.update(文=f"✓ {len(結果['頁們'])} 頁", 色=介.OK)

    def 加入():
        if 忙["是"]:
            return
        起點 = os.path.dirname(檔們[-1]) if 檔們 else ""
        選 = filedialog.askopenfilenames(parent=介.root, title="選擇客戶的掃描 PDF（可以多選）", initialdir=起點,
                                       filetypes=[("PDF 或圖片", "*.pdf *.png *.jpg *.jpeg"), ("所有檔案", "*.*")])
        for p in 選 or ():
            if p not in 檔們:
                檔們.append(p)
        畫(); 辨識()

    def 移除(p):
        if 忙["是"]:
            return
        檔們.remove(p)
        畫(); 辨識()

    def 匯出():
        if not 結果["頁們"]:
            介.messagebox.showinfo("匯出辨識結果", "還沒有辨識結果", parent=介.root)
            return
        p = filedialog.asksaveasfilename(parent=介.root, title="辨識結果存成", defaultextension=".xlsx",
                                         initialfile="辨識結果.xlsx", filetypes=[("Excel 檔案", "*.xlsx")])
        if p:
            匯出辨識結果(結果["頁們"], p)
            介.messagebox.showinfo("匯出辨識結果", f"已存成：{p}", parent=介.root)

    def 收集(設):
        設["掃描檔們"] = list(檔們)
        設["尾牙月份"] = int(尾牙月.get().replace("月", "") or 2)

    def 狀態():
        if 忙["是"]:
            return {"檔案": ("辨識中…", 介.WARN)}
        return {"檔案": (狀["文"], 狀["色"])} if 狀["文"] else {}
    畫()
    def 產生前確認(設):
        if 忙["是"]:
            介.messagebox.showinfo("還在辨識", "掃描檔還在辨識，等進度跑完再按「產生」。", parent=介.root)
            return False
        return True
    return {"收集": 收集, "載入": lambda: (畫(), 辨識()) if not 結果["頁們"] and 檔們 else 畫(), "狀態": 狀態,
            "產生前確認": 產生前確認}