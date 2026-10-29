# -*- coding: utf-8 -*-
"""
清冊共用：薪資印領清冊（共用核心＋客戶模組）
==========================================
每個客戶的薪資檔長得都不一樣，所以分成兩部分：
  ・共用核心（這個檔案）：清冊樣板、填清冊、驗算（個人加總 vs 表上合計）、匯入檢查、
    人員名冊（編號／身分證／地址以名冊為準＋核對）、科目分類、視窗。所有客戶共用，不用改。
  ・客戶模組（旁邊「客戶模組」資料夾，一個 .py 一種讀法）：這個客戶的薪資檔要怎麼讀。
      通用.py ─ 一個 Excel、每月一張表，在畫面上設定欄位就好（格式簡單的客戶用這個）
      七賢.py ─ 七賢：扣憑＋自提、獎金可另放月份
      澤利.py ─ 澤利：應付薪資＋員工酬勞、依職稱分會計科目
    新客戶格式很特別（每月一個檔、PDF、好幾欄加總、分店…）→ 複製一個模組改「讀取」就好。

客戶模組裡可以有（都可以省略）：
  名稱、說明          畫面上顯示的名字和說明
  預設設定            新增這種客戶時的設定（dict）
  預設客戶            第一次執行時自動建好的客戶名稱
  頁面(介, 設)        左邊選單多幾頁（介 裡有做卡片、輸入框、下拉的小工具）
  整理設定(設)        產生前把自己的設定換成共用核心的設定
  讀取(設, 年, wb, wbf, 清冊欄) → (表們, 缺表, 警告)
                      自己讀薪資檔：每張表做一個 資料來源，用 .加(姓名, 總額, 自提, 稅…) 放人進去

使用方式（需先安裝：pip install openpyxl xlrd）
  雙擊或執行  python 清冊共用.py → 左上角選客戶，按「產生」
  指令：python 清冊共用.py 客戶名稱 [輸出資料夾]（用該客戶存好的設定）
  每個客戶的設定存在「客戶設定」資料夾（一個客戶一個 .json）

注意：薪資檔裡的公式，工具讀的是 Excel 存檔時算好的值，薪資檔請先用 Excel 開啟並存檔過一次。
"""
import copy
import difflib
import json
import os
import re
import sys
from collections import Counter, OrderedDict
from decimal import Decimal, ROUND_HALF_UP

try:
    import openpyxl
    from openpyxl.cell.cell import MergedCell
    from openpyxl.formula.translate import Translator
    from openpyxl.styles import Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter, column_index_from_string
except ImportError:
    sys.exit("找不到 openpyxl，請先執行：pip install openpyxl")

# ============================== 認標題用的字（一般不用改）==============================
姓名標題 = ["姓名", "醫師別", "醫師姓名", "員工姓名", "所得人姓名", "所得人"]
編號標題 = ["員工編號", "編號", "員編", "工號"]
身分證關鍵字 = ["身分證", "身份證", "統一編號", "統編", "居留證"]
# 獎金表的金額欄：先找完全一樣的標題（由前往後），再找「含有」的
金額標題_優先 = ["應發金額", "應發放金額", "應付年終", "給付金額", "應領金額", "金額", "獎金"]
金額標題_次要 = ["合計"]
金額關鍵字 = ["禮金", "獎金", "金額"]
金額排除字 = ["實發", "實付", "實領", "扣", "稅", "健保", "級距"]
扣繳排除字 = ["級距", "實發", "實付", "合計", "總計"]
非人名 = {"小計", "合計", "總計", "總合計", "TOTAL", "總額"}
中文月 = ["一月", "二月", "三月", "四月", "五月", "六月", "七月", "八月", "九月", "十月", "十一月", "十二月"]

# 固定的處理方式
只列有資料的人 = True     # 只列今年有金額的人
零不填 = True            # 金額 0 的格子留空白
四捨五入到元 = True       # 有小數的金額四捨五入到整數

# 執行時依客戶設定換掉（月薪表的認法）
月薪金額詞 = ["扣憑"]


# ================================================================ 小工具
def 清字(v):
    """去掉空白、換行：「 應發\\n金額 」→「應發金額」"""
    if v is None:
        return ""
    return re.sub(r"\s+", "", str(v).replace("　", ""))


def 是數字(v):
    if isinstance(v, bool) or v is None:
        return False
    if isinstance(v, (int, float)):
        return True
    try:
        float(str(v).replace(",", ""))
        return True
    except ValueError:
        return False


def 數字(v):
    if not 是數字(v):
        return 0
    x = float(v) if isinstance(v, (int, float)) else float(str(v).replace(",", ""))
    if 四捨五入到元:
        return int(Decimal(repr(x)).quantize(Decimal("1"), rounding=ROUND_HALF_UP))
    return int(x) if x.is_integer() else round(x, 2)


def 是人名(n):
    return bool(n) and n.upper() not in 非人名 and not 是數字(n) and len(n) <= 10


def 編號鍵(v):
    t = 清字(v)
    if not t:
        return ""
    return str(int(t)) if t.isdigit() else t.upper()


def 拆詞(文字):
    """「應付薪資＋員工酬勞」→ ["應付薪資", "員工酬勞"]"""
    if isinstance(文字, (list, tuple)):
        return [清字(x) for x in 文字 if 清字(x)]
    return [清字(x) for x in re.split(r"[+＋、,，;；]", 文字 or "") if 清字(x)]


def 拆對照(文字):
    """「5300=00、5401=01」→ [("5300", "00"), ("5401", "01")]"""
    結果 = []
    for x in re.split(r"[、,，;；\n]+", 文字 or ""):
        if "=" in x:
            a, b = x.split("=", 1)
            if a.strip():
                結果.append((a.strip(), b.strip()))
    return 結果


# 表名：把設定裡的 {年度} {前一年} {月} {月2} 換成實際的字
#   "{年度}年{月}月" → "115年3月"；"{年度}{月2}" → "11403"；"{前一年}員工年終" → "114員工年終"
def 表名(模式, 年, 月=None):
    s = 模式.replace("{年度}", str(年)).replace("{前一年}", str(年 - 1))
    if 月 is not None:
        s = s.replace("{月2}", f"{月:02d}").replace("{月}", str(月))
    return s


# 轉模式：視窗裡選了「114員工年終」，存成 "{前一年}員工年終"，明年自動變 115員工年終
def 轉模式(名, 年):
    m = re.match(r"^(\d{2,4})(.*)$", 名.strip())
    if m and int(m.group(1)) == 年:
        return "{年度}" + m.group(2)
    if m and int(m.group(1)) == 年 - 1:
        return "{前一年}" + m.group(2)
    return 名.strip()


def 找工作表(wb, 名稱):
    for ws in wb.worksheets:
        if 清字(ws.title) == 清字(名稱):
            return ws
    return None


# ================================================================ 讀表頭
class 表:
    """一張工作表裡的一個資料表（同一張工作表可能並排兩個表，例如端午節的員工表和醫師表）"""

    def __init__(self, ws, 表頭列, 姓名欄, 標題):
        self.ws, self.表頭列, self.姓名欄, self.標題 = ws, 表頭列, 姓名欄, 標題   # 標題：{文字: 欄}
        self.工作表 = ws.title
        self.姓名標題 = next(t for t, c in 標題.items() if c == 姓名欄)
        self.資料起始列 = 表頭列 + 1
        self.編號欄 = next((標題[t] for t in 編號標題 if t in 標題), None)
        self.月薪 = any(any(k in t for k in 月薪金額詞) for t in 標題)
        self.總額 = []          # 薪資總額＝這幾個標題相加
        self.自提 = None
        self.扣繳 = []
        self.欄位 = None        # 填進清冊哪一欄
        self.類型 = "獎金"

    @property
    def 金額(self):
        return self.總額[0] if self.總額 else None

    @property
    def 顯示名稱(self):
        return self.工作表.strip() + ("" if self.姓名標題 == "姓名" else f"（{self.姓名標題}）")


def 合併範圍(ws):
    m = {}
    for rg in ws.merged_cells.ranges:
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1):
                m[(r, c)] = rg
    return m


def 標題字(cell):
    """在 Excel 打 10%，存起來是數字 0.1 → 還原成文字「10%」"""
    v = cell.value
    if isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in (cell.number_format or ""):
        return f"{v * 100:g}%"
    return 清字(v)


# 找資料表：前 10 列找「姓名」標題，往左右延伸收集同一列的標題；
#   橫向合併的大標題（例如「其他扣款」底下再分小標題）改讀下一列
def 找資料表(ws):
    合併 = 合併範圍(ws)
    結果 = []
    最大欄 = min(ws.max_column, 60)

    def 有字(r, c):
        v = ws.cell(r, c).value
        if (r, c) in 合併 and (合併[(r, c)].min_row, 合併[(r, c)].min_col) != (r, c):
            return True
        if 標題字(ws.cell(r, c)).endswith("%") and 是數字(v):
            return True
        return v is not None and not 是數字(v) and 清字(v) != ""

    for r in range(1, min(10, ws.max_row) + 1):
        for c in range(1, 最大欄 + 1):
            if 清字(ws.cell(r, c).value) not in 姓名標題:
                continue
            左 = c
            while 左 > 1:
                if 有字(r, 左 - 1):
                    左 -= 1
                elif 左 > 2 and 有字(r, 左 - 2) and not 是數字(ws.cell(r, 左 - 1).value):
                    左 -= 2
                else:
                    break
            右 = c
            while 右 < 最大欄:
                if 有字(r, 右 + 1):
                    右 += 1
                elif 右 + 2 <= 最大欄 and 有字(r, 右 + 2) and not 是數字(ws.cell(r, 右 + 1).value):
                    右 += 2
                else:
                    break
            標題, 子標題 = {}, False
            for cc in range(左, 右 + 1):
                rg = 合併.get((r, cc))
                t = 標題字(ws.cell(r, cc))
                if rg and rg.max_col > rg.min_col and rg.min_row == r:
                    t2 = 標題字(ws.cell(r + 1, cc))
                    if t2:
                        標題.setdefault(t2, cc)
                        子標題 = True
                        continue
                if t:
                    標題.setdefault(t, cc)
            t = 表(ws, r, c, 標題)
            if 子標題:
                t.資料起始列 = r + 2
            結果.append(t)
    return 結果


# 找標題：用設定的字找表裡的標題。先找完全一樣的，沒有再找「含有」的
def 找標題(t, 詞, 排除=()):
    詞 = 清字(詞)
    if not 詞:
        return None
    if 詞 in t.標題:
        return 詞
    return next((標 for 標 in t.標題 if 詞 in 標 and not any(x in 標 for x in 排除)), None)


def 選金額欄(t):
    h = t.標題
    for k in 金額標題_優先 + 金額標題_次要:
        if k in h:
            return k
    for k in 金額關鍵字:
        for 標 in h:
            if k in 標 and not any(x in 標 for x in 金額排除字):
                return 標
    return None


def 選扣繳欄(t, 詞):
    return [標 for 標 in t.標題 if any(清字(k) and 清字(k) in 標 for k in 詞)
            and not any(x in 標 for x in 扣繳排除字)]


# 設定欄位：依客戶設定決定這個表的 薪資總額／自提／扣繳 是哪幾個標題
#   月薪格式：薪資總額＝設定的幾欄相加、自提＝設定的那欄、扣繳＝設定的幾欄
#   獎金格式：金額欄自動找（或設定指定）、扣繳＝標題含「獎金扣繳詞」的欄
def 設定欄位(t, 類型, 設):
    t.類型 = 類型
    if 類型 == "月薪":
        t.總額 = [x for x in (找標題(t, k) for k in 拆詞(設["薪資總額"])) if x]
        t.自提 = next((x for x in (找標題(t, k) for k in 拆詞(設["退休自提"])) if x), None)
        t.扣繳 = 選扣繳欄(t, 拆詞(設["扣繳稅額"]))
    else:
        t.總額 = [x for x in [選金額欄(t)] if x]
        t.自提 = None
        t.扣繳 = 選扣繳欄(t, 拆詞(設.get("獎金扣繳詞", "")))
    return bool(t.總額)


# 月薪表的資料表：有「薪資總額」第一個字（例如 扣憑、應付薪資）的那個表
def 月薪資料表(ws, 設):
    詞 = 拆詞(設["薪資總額"])
    for t in 找資料表(ws):
        if 詞 and 找標題(t, 詞[0]):
            設定欄位(t, "月薪", 設)
            return t
    return None


# 選表：依客戶設定找出要讀的表。回傳 (表們, 缺少的工作表, 警告)
def 選表(wb, 設, 年, 清冊欄):
    表們, 缺, 警告 = [], [], []
    缺自提, 缺扣繳 = [], []
    清單 = [(表名(設["月薪表"], 年, m), 中文月[m - 1], "月薪", False) for m in range(1, 13)]
    清單 += [(表名(x[0], 年), 清字(x[1]), x[2] if len(x) > 2 and x[2] else "獎金", True) for x in 設.get("其他表", [])]
    for 名, 欄位, 類型, 必要 in 清單:
        if 欄位 not in 清冊欄:
            raise SystemExit(f"「{名}」要填入「{欄位}」，但樣板沒有這一欄")
        ws = 找工作表(wb, 名)
        if ws is None:
            if 必要:
                缺.append(名)
            continue
        if 類型 == "月薪":
            t = 月薪資料表(ws, 設)
            if t is None:
                raise SystemExit(f"「{ws.title.strip()}」找不到姓名欄，或找不到薪資總額的欄位（{設['薪資總額']}），"
                                 f"請檢查客戶設定的「薪資總額」")
            t.欄位 = 欄位
            if 拆詞(設["退休自提"]) and not t.自提:
                缺自提.append(ws.title.strip())
            if 拆詞(設["扣繳稅額"]) and not t.扣繳:
                缺扣繳.append(ws.title.strip())
            表們.append(t)
            continue
        找到, 有姓名 = 0, []
        for t in 找資料表(ws):
            有姓名.append(t)
            if not 設定欄位(t, "獎金", 設):
                continue
            t.欄位 = 欄位
            表們.append(t)
            找到 += 1
        if not 有姓名:
            raise SystemExit(f"「{ws.title.strip()}」前 10 列找不到姓名標題（{'、'.join(姓名標題)}），請檢查標題")
        if not 找到:
            raise SystemExit(f"「{ws.title.strip()}」找不到金額欄（標題要是 應發金額、金額、獎金、合計，"
                             f"或含 禮金／獎金／金額），請檢查標題")
    if 缺自提:
        警告.append(f"月薪表找不到退休自提欄（{設['退休自提']}），自提沒有填：{'、'.join(缺自提)}")
    if 缺扣繳:
        警告.append(f"月薪表找不到扣繳欄（{設['扣繳稅額']}），扣繳稅額沒有填：{'、'.join(缺扣繳)}")
    return 表們, 缺, 警告


# ================================================================ 讀資料
# 讀名冊：名稱含「名冊」的工作表 → {姓名: 編號、地址、身分證號}、{編號: 姓名}。沒有名冊就回傳空的
def 讀名冊(wb):
    for ws in wb.worksheets:
        if "名冊" not in ws.title:
            continue
        for r in range(1, 10):
            h = {清字(ws.cell(r, c).value): c for c in range(1, ws.max_column + 1) if 清字(ws.cell(r, c).value)}
            名 = next((h[k] for k in 姓名標題 if k in h), None)
            if not 名:
                continue
            編 = next((h[k] for k in 編號標題 if k in h), None)
            址 = next((c for k, c in h.items() if "地址" in k), None)
            證 = next((c for k, c in h.items() if any(x in k for x in 身分證關鍵字)), None)
            名冊, 編號 = OrderedDict(), {}
            for rr in range(r + 1, ws.max_row + 1):
                n = 清字(ws.cell(rr, 名).value)
                if not 是人名(n) or n in 名冊:
                    continue
                名冊[n] = {"編號": ws.cell(rr, 編).value if 編 else None,
                          "地址": ws.cell(rr, 址).value if 址 else None,
                          "身分證號": 清字(ws.cell(rr, 證).value).upper() or None if 證 else None}
                k = 編號鍵(名冊[n]["編號"])
                if k:
                    編號[k] = n
            return ws.title, 名冊, 編號
    return None, OrderedDict(), {}


# 讀身分證（備用）：名冊沒有的人，到任何有「姓名」和「身分證」標題的工作表找
def 讀身分證(wb):
    證 = {}
    for ws in wb.worksheets:
        for hr in range(1, 11):
            h = {清字(ws.cell(hr, c).value): c for c in range(1, min(ws.max_column, 60) + 1) if 清字(ws.cell(hr, c).value)}
            名欄 = next((c for k, c in h.items() if k in 姓名標題), None)
            證欄 = next((c for k, c in h.items() if any(x in k for x in 身分證關鍵字)), None)
            if 名欄 and 證欄:
                for r in range(hr + 1, ws.max_row + 1):
                    n, v = 清字(ws.cell(r, 名欄).value), 清字(ws.cell(r, 證欄).value).upper()
                    if 是人名(n) and len(v) >= 8:
                        證[n] = v
                break
    return 證


# 找合計列：姓名欄寫「合計」，或姓名欄空白但金額欄是 =SUM(...) 的那一列
def 找合計列(ws, wsf, t):
    金欄 = t.標題[t.金額] if t.金額 else None
    for r in range(t.資料起始列, ws.max_row + 1):
        n = 清字(ws.cell(r, t.姓名欄).value)
        if n.upper() in 非人名:
            return r
        if 金欄 and not n:
            f = wsf.cell(r, 金欄).value
            if isinstance(f, str) and f.replace(" ", "").upper().startswith("=SUM("):
                return r
    return None


# 讀表資料：一列一列讀，每個人整理成 {名, 編, 金:{總額,自提,稅}, 執, 職, 證}
#   順便核對：每一欄的個人加總 vs 表上合計列
def 讀表資料(t, wbf, 警告):
    ws, h = t.ws, t.標題
    wf = wbf[ws.title]
    cn = t.姓名欄
    總欄 = [h[x] for x in t.總額]
    cs = h.get(t.自提) if t.自提 else None
    ct = [h[x] for x in t.扣繳 if x in h]
    職欄 = h.get("職稱")
    證欄 = next((c for k, c in h.items() if any(x in k for x in 身分證關鍵字)), None)
    欄們 = sorted(set(h.values()))
    左, 右 = 欄們[0], 欄們[-1]
    合計列 = 找合計列(ws, wf, t)
    列們, 無值, 非數字 = [], [], []
    for r in range(t.資料起始列, ws.max_row + 1):
        if r == 合計列:
            continue
        n = 清字(ws.cell(r, cn).value)
        if not 是人名(n):
            continue
        av = ws.cell(r, 總欄[0]).value
        if av is None:
            if str(wf.cell(r, 總欄[0]).value or "").startswith("="):
                無值.append(f"{n}(第{r}列)")
            continue
        if not 是數字(av):
            非數字.append(f"{n}(第{r}列)「{av}」")
            continue
        自 = 數字(ws.cell(r, cs).value) if cs else 0
        金 = {"總額": sum(數字(ws.cell(r, c).value) for c in 總欄), "自提": 自,
              "稅": sum(數字(ws.cell(r, c).value) for c in ct)}
        列們.append({"名": n, "編": 編號鍵(ws.cell(r, t.編號欄).value) if t.編號欄 else "", "金": 金,
                   "執": any("執行業務" in 清字(ws.cell(r, c).value) for c in range(左, 右 + 1)),
                   "職": 清字(ws.cell(r, 職欄).value) if 職欄 else "",
                   "證": 清字(ws.cell(r, 證欄).value).upper() if 證欄 else ""})
    def 列舉(x):
        return "、".join(x[:5]) + (f" 等 {len(x)} 列" if len(x) > 5 else "")
    if 無值:
        警告.append(f"「{t.顯示名稱}」{t.金額} 是公式但沒有計算值，未填入：{列舉(無值)}。請用 Excel 開啟薪資檔、存檔後再執行")
    if 非數字:
        警告.append(f"「{t.顯示名稱}」{t.金額} 不是數字，未填入：{列舉(非數字)}")
    t.合計列, t.核對 = 合計列, []
    if 合計列:                               # 每一欄：個人加總 vs 表上合計
        for 標 in t.總額 + ([t.自提] if t.自提 else []) + t.扣繳:
            c = h[標]
            表上 = ws.cell(合計列, c).value
            if not 是數字(表上):
                continue
            加總 = sum(數字(ws.cell(r, c).value) for r in range(t.資料起始列, 合計列)
                     if 是人名(清字(ws.cell(r, cn).value)))
            t.核對.append((標, 加總, 數字(表上)))
            if abs(加總 - 數字(表上)) > 1:
                警告.append(f"「{t.顯示名稱}」{標} 個人加總 {加總:,} ≠ 表上合計 {數字(表上):,}（第 {合計列} 列），請確認")
    return 列們


# ================================================================ 寫入（合併儲存格安全）
def 寫(ws, r, c, v, 清除=False):
    """寫進一格；如果是合併儲存格裡不是左上角的格子，改寫到左上角（清除的話就跳過）"""
    格 = ws.cell(r, c)
    if isinstance(格, MergedCell):
        rg = next((m for m in ws.merged_cells.ranges
                   if m.min_row <= r <= m.max_row and m.min_col <= c <= m.max_col), None)
        if rg is None or 清除:
            return
        格 = ws.cell(rg.min_row, rg.min_col)
    格.value = v


# ================================================================ 清冊樣板
def 清冊欄位表(ws):
    """在樣板裡找「明細」那一格 → (表頭列, 明細欄, {欄名: 欄})"""
    for r in range(1, 10):
        for c in range(1, ws.max_column + 1):
            if 清字(ws.cell(r, c).value) == "明細":
                欄 = OrderedDict()
                for cc in range(1, ws.max_column + 1):
                    t = 清字(ws.cell(r, cc).value)
                    if t:
                        欄[t] = cc
                return r, c, 欄
    raise SystemExit("樣板找不到「明細」欄，無法判斷表頭在哪一列")


def 樣板工作表(wb):
    for ws in wb.worksheets:
        for r in range(1, 10):
            if any(清字(ws.cell(r, c).value) == "明細" for c in range(1, min(ws.max_column, 40) + 1)):
                return ws
    raise SystemExit("樣板裡找不到有「明細」表頭的工作表")


def 移欄公式(公式, 位置):
    def f(m):
        欄 = column_index_from_string(m.group(2))
        if 欄 >= 位置:
            欄 += 1
        return m.group(1) + get_column_letter(欄) + m.group(3) + m.group(4)
    return re.sub(r"(\$?)([A-Z]{1,3})(\$?)(\d+)", f, 公式)


# 插入欄：清冊沒有設定裡的欄位（例如中秋獎金）→ 在「本期年終」前面插入一欄，公式、格式、合併一起移
def 插入欄(ws, 位置, 名稱, 表頭列):
    最大列, 最大欄 = ws.max_row, ws.max_column
    合併們 = [(m.min_row, m.min_col, m.max_row, m.max_col) for m in ws.merged_cells.ranges]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))
    for r in range(1, 最大列 + 1):
        for c in range(1, 最大欄 + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.startswith("="):
                ws.cell(r, c).value = 移欄公式(v, 位置)
    for c in range(最大欄, 位置 - 1, -1):
        for r in range(1, 最大列 + 1):
            來源, 目的 = ws.cell(r, c), ws.cell(r, c + 1)
            目的.value = 來源.value
            目的._style = copy.copy(來源._style)
        寬 = ws.column_dimensions[get_column_letter(c)].width
        if 寬:
            ws.column_dimensions[get_column_letter(c + 1)].width = 寬
    左 = 位置 - 1
    for r in range(1, 最大列 + 1):
        左格, 新格 = ws.cell(r, 左), ws.cell(r, 位置)
        新格._style = copy.copy(左格._style)
        v = 左格.value
        if isinstance(v, str) and v.startswith("=") and r > 表頭列:
            新格.value = Translator(v, origin=左格.coordinate).translate_formula(新格.coordinate)
        else:
            新格.value = None
    ws.cell(表頭列, 位置).value = 名稱
    寬 = ws.column_dimensions[get_column_letter(左)].width
    if 寬:
        ws.column_dimensions[get_column_letter(位置)].width = 寬
    for a, b, c, d in 合併們:
        if b >= 位置:
            b, d = b + 1, d + 1
        elif d >= 位置:
            d += 1
        ws.merge_cells(start_row=a, start_column=b, end_row=c, end_column=d)


def 換列號(公式, 對照):
    def f(m):
        return m.group(1) + str(對照(int(m.group(2))))
    return re.sub(r"(\$?[A-Z]{1,3}\$?)(\d+)(?![\d(])", f, 公式)


# 樣板：讀懂清冊樣板，依人數重新長出清冊
#   ・每人一組幾列、哪一列放總額／自提／稅（看「明細」欄的字）
#   ・空白樣板（表頭下面沒有任何一組）→ 自動做一組：薪資總額、退休自提（有扣繳才加扣繳稅額）
#   ・最下面的合計列照搬，公式範圍自動改
class 樣板:
    def __init__(self, ws, 要稅=False):
        self.ws = ws
        self.表頭列, self.明細欄, self.欄 = 清冊欄位表(ws)
        self.欄數 = max(ws.max_column, max(self.欄.values()))
        F = self.明細欄
        lab = lambda r: 清字(ws.cell(r, F).value)
        r = self.表頭列 + 1
        while r <= ws.max_row and not lab(r):
            r += 1
        self.首列 = self.表頭列 + 1
        沒標籤 = r > ws.max_row
        if 沒標籤:
            # 表頭下面沒有「明細」字：如果表頭下面已經畫好一組（有合併或框線），就照那組格式，自己補上明細字；
            #   什麼都沒有就自己做一組
            h1 = self.表頭列 + 1
            高 = max([m.max_row - h1 + 1 for m in ws.merged_cells.ranges if m.min_row == h1 and m.max_row > h1] or [0])
            有格式 = any(ws.cell(h1, c).has_style and ws.cell(h1, c).border.left.style for c in range(1, self.欄數 + 1))
            if not 高 and not 有格式:
                self._自己做一組(要稅)
                return
            標們 = ["薪資總額", "退休自提"] + (["扣繳稅額"] if 要稅 else [])
            高 = 高 or len(標們)
            for i in range(高):
                if i < len(標們):
                    寫(ws, h1 + i, F, 標們[i])
            r = h1
        self.首列 = r
        首標 = lab(r)
        組 = [r]
        rr = r + 1
        while rr <= ws.max_row and lab(rr) != 首標 and (not 沒標籤 or rr < r + 高):
            組.append(rr)
            if "蓋章" in lab(rr):
                break
            rr += 1
        self.組列, self.組長 = 組, len(組)
        self.角色 = {}
        for i, x in enumerate(組):
            t = lab(x)
            if "自提" in t:
                self.角色["自提"] = i
            elif "扣繳" in t or "稅" in t:
                self.角色["稅"] = i
            elif "蓋章" not in t and "總額" not in self.角色:
                self.角色["總額"] = i
        if "總額" not in self.角色:
            raise SystemExit("樣板第一組裡找不到薪資總額那一列")
        末 = r
        while not 沒標籤 and lab(末 + self.組長) == 首標:
            末 += self.組長
        self.舊末列 = 末 + self.組長 - 1

        def snap(x, 全留=False):                 # 每組只留「明細」字和公式（樣板裡範例人的資料不要複製）
            列 = []
            for c in range(1, self.欄數 + 1):
                v = ws.cell(x, c).value
                if not 全留 and c != F and not (isinstance(v, str) and v.startswith("=")):
                    v = None
                列.append((v, copy.copy(ws.cell(x, c)._style)))
            return 列, ws.row_dimensions[x].height
        self.組樣 = [snap(x) for x in 組]
        self.組合併 = [(m.min_row - r, m.min_col, m.max_row - r, m.max_col) for m in ws.merged_cells.ranges
                     if m.min_row >= r and m.max_row <= 組[-1]]
        # 基本資料要不要每一列都填：樣板範例在第二列以後也有寫（而且沒合併）的欄才每列都填
        直合併 = {c for a, b, cc, d in self.組合併 if cc > a for c in range(b, d + 1)}
        self.每列填 = {c for c in range(1, self.欄數 + 1) if c not in 直合併
                     and any(ws.cell(x, c).value not in (None, "") for x in 組[1:]) and c != F}
        self.尾 = [(x - self.舊末列, *snap(x, True)) for x in range(self.舊末列 + 1, ws.max_row + 1)]
        self.尾合併 = [(m.min_row - self.舊末列, m.min_col, m.max_row - self.舊末列, m.max_col)
                     for m in ws.merged_cells.ranges if m.min_row > self.舊末列]

    def _自己做一組(self, 要稅):
        ws = self.ws
        細 = Side(style="thin")
        框 = Border(left=細, right=細, top=細, bottom=細)
        標們 = ["薪資總額", "退休自提"] + (["扣繳稅額"] if 要稅 else [])
        self.組樣 = []
        for 標 in 標們:
            列 = []
            for c in range(1, self.欄數 + 1):
                格 = ws.cell(self.表頭列 + 50, c)       # 一格沒用過的格子當底，再加框線
                格.border = 框
                列.append((標 if c == self.明細欄 else None, copy.copy(格._style)))
                格.border = Border()
            self.組樣.append((列, None))
        self.組列 = list(range(self.首列, self.首列 + len(標們)))
        self.組長 = len(標們)
        self.角色 = {"總額": 0, "自提": 1, **({"稅": 2} if 要稅 else {})}
        self.組合併, self.尾, self.尾合併 = [], [], []
        self.舊末列 = self.表頭列
        self.每列填 = {c for c in range(1, self.欄數 + 1) if c != self.明細欄}

    # 加獎金列：在每組「薪資總額」那一列下面，多加幾列給「選了月份的獎金」（例如春節獎金、發財獎），
    #   格式照薪資總額那列複製，明細欄寫獎金名稱；退休金自提、扣繳稅額、蓋章往下移
    def 加獎金列(self, 名們):
        if not 名們:
            return
        k = len(名們)
        i0 = self.角色["總額"]
        p = i0 + 1                                  # 插在薪資總額的下一列
        組首 = self.組列[0]
        舊長 = self.組長

        def 移(公式, 對照):
            return 換列號(公式, 對照) if isinstance(公式, str) and 公式.startswith("=") else 公式
        往下 = lambda x: x + k if 組首 + p <= x <= 組首 + 舊長 - 1 else x
        新樣 = []
        for i, (列, 高) in enumerate(self.組樣):
            if i == p:
                for j, 名 in enumerate(名們):
                    到 = lambda x, j=j: 組首 + p + j if x == 組首 + i0 else x
                    新列 = [((名 if c == self.明細欄 else 移(v, 到)), st) for c, (v, st) in enumerate(self.組樣[i0][0], 1)]
                    新樣.append((新列, self.組樣[i0][1]))
            新樣.append(([(移(v, 往下), st) for v, st in 列], 高))
        if p >= len(self.組樣):
            for j, 名 in enumerate(名們):
                到 = lambda x, j=j: 組首 + p + j if x == 組首 + i0 else x
                新樣.append(([((名 if c == self.明細欄 else 移(v, 到)), st) for c, (v, st) in enumerate(self.組樣[i0][0], 1)],
                            self.組樣[i0][1]))
        self.組樣 = 新樣
        self.組長 = 舊長 + k
        self.組列 = list(range(組首, 組首 + self.組長))
        self.組合併 = [((a + k, b, c + k, d) if a >= p else ((a, b, c + k, d) if c >= p else (a, b, c, d)))
                     for a, b, c, d in self.組合併]
        self.角色 = {角: (i + k if i >= p else i) for 角, i in self.角色.items()}
        for j, 名 in enumerate(名們):
            self.角色["獎:" + 名] = p + j

    def 重建(self, 人數):
        ws = self.ws
        for m in list(ws.merged_cells.ranges):
            if m.max_row >= self.首列:
                ws.unmerge_cells(str(m))
        無格式 = copy.copy(ws.cell(ws.max_row + 50, 1)._style)
        for r in range(self.首列, max(ws.max_row, self.首列 + self.組長 * 人數 + len(self.尾)) + 1):
            for c in range(1, self.欄數 + 1):
                ws.cell(r, c).value = None
                ws.cell(r, c)._style = copy.copy(無格式)
        新末 = self.首列 + self.組長 * 人數 - 1
        差 = 新末 - self.舊末列
        組首, 組尾 = self.組列[0], self.組列[-1]

        def 貼(r, 列, 高, 對照):
            for c, (v, st) in enumerate(列, 1):
                cell = ws.cell(r, c)
                cell._style = copy.copy(st)
                if isinstance(v, str) and v.startswith("="):
                    v = 換列號(v, 對照)
                cell.value = v
            ws.row_dimensions[r].height = 高

        for k in range(人數):
            r0 = self.首列 + self.組長 * k
            位移 = r0 - 組首
            對照 = lambda x, d=位移: x + d if 組首 <= x <= 組尾 else x
            for i, (列, 高) in enumerate(self.組樣):
                貼(r0 + i, 列, 高, 對照)
            for a, b, c, d in self.組合併:
                ws.merge_cells(start_row=r0 + a, start_column=b, end_row=r0 + c, end_column=d)
        對照尾 = lambda x: 新末 if x == self.舊末列 else (x + 差 if x > self.舊末列 else x)
        for 偏, 列, 高 in self.尾:
            貼(新末 + 偏, 列, 高, 對照尾)
        for a, b, c, d in self.尾合併:
            ws.merge_cells(start_row=新末 + a, start_column=b, end_row=新末 + c, end_column=d)

    # 填：一個人一組。基本資料（編號、身分證號、姓名、地址、科目）、金額依角色填
    def 填(self, 人們, 資料, 基本, 顯示名, 編號們=None, 科目們=None, 科目欄名="", 月份們=None):
        ws = self.ws
        月份們 = 月份們 or {}

        def 加(r, c, v):                          # 同一格可能有好幾筆（例如這個月的薪資扣繳＋獎金扣繳）→ 加起來
            舊 = ws.cell(r, c).value
            寫(ws, r, c, (舊 if isinstance(舊, (int, float)) else 0) + v)
        self.重建(max(len(人們), 1))
        證欄 = next((c for k, c in self.欄.items() if any(x in k for x in 身分證關鍵字)), None)
        名欄 = next((c for k, c in self.欄.items() if "姓名" in k), None)
        科欄 = next((c for k, c in self.欄.items() if 科目欄名 and 清字(科目欄名) in k), None) if 科目們 else None

        def 放(r0, c, v, 文字=False):
            if not c or v in (None, ""):
                return
            for i in range(self.組長):
                if i == 0 or c in self.每列填:
                    寫(ws, r0 + i, c, v)
                    if 文字 and not isinstance(ws.cell(r0 + i, c), MergedCell) and ws.cell(r0 + i, c).number_format == "General":
                        ws.cell(r0 + i, c).number_format = "@"     # 編號、身分證號存成文字（0 開頭不會不見）

        for k, 人 in enumerate(人們):
            r0 = self.首列 + self.組長 * k
            info = 基本.get(人, {})
            放(r0, self.欄.get("編號"), (編號們 or {}).get(人, info.get("編號")), True)
            放(r0, 名欄, 顯示名.get(人, 人))
            放(r0, self.欄.get("地址"), info.get("地址"))
            放(r0, 證欄, info.get("身分證號"), True)
            if 科欄:
                科 = (科目們 or {}).get(人, "")
                放(r0, 科欄, int(科) if str(科).isdigit() else 科)
            項們 = list(資料.get(人, {}).items())
            for 欄名, v in [x for x in 項們 if x[0] not in 月份們] + [x for x in 項們 if x[0] in 月份們]:
                c = self.欄[欄名]
                月 = 月份們.get(欄名)
                for 角, 鍵 in (("總額", "總額"), ("自提", "自提"), ("稅", "稅")):
                    if 角 in self.角色 and (v[鍵] or not 零不填):
                        寫(ws, r0 + self.角色[角], c, v[鍵])
                        if 月 and 角 == "稅" and v[鍵]:   # 選了月份的獎金：扣繳也加進那個月的扣繳稅額
                            加(r0 + self.角色[角], self.欄[中文月[月 - 1]], v[鍵])
                if 月 and v["總額"] and ("獎:" + 欄名) in self.角色:   # 獎金自己那一列：放在選的月份
                    加(r0 + self.角色["獎:" + 欄名], self.欄[中文月[月 - 1]], v["總額"])




# ================================================================ 科目分類（選用）
# 插入列：在 at 這列前面插入 n 列，公式裡的列號、合併儲存格、列高一起往下移（別張表參照到的也會改）
def 插入列(wb, ws, at, n):
    from openpyxl.formula.tokenizer import Tokenizer
    參照 = re.compile(r"(\$?)([A-Za-z]{1,3})(\$?)(\d+)")

    def 移(公式, 本表):
        try:
            項 = Tokenizer(公式).items
        except Exception:
            return 公式
        結果 = "="
        for x in 項:
            v = x.value
            if x.type == "OPERAND" and x.subtype == "RANGE":
                表_, 位 = (v.rsplit("!", 1) if "!" in v else (None, v))
                目標 = 表_.strip("'") if 表_ else 本表
                if 目標 == ws.title:
                    位 = 參照.sub(lambda m: m.group(1) + m.group(2) + m.group(3)
                                + str(int(m.group(4)) + n if int(m.group(4)) >= at else m.group(4)), 位)
                    v = (表_ + "!" if 表_ else "") + 位
            結果 += v
        return 結果
    合併 = [m for m in ws.merged_cells.ranges if m.max_row >= at]
    for m in 合併:
        ws.unmerge_cells(str(m))
    高 = {r: d.height for r, d in ws.row_dimensions.items() if r >= at and d.height}
    ws.insert_rows(at, n)
    for r in sorted(高, reverse=True):
        ws.row_dimensions[r + n].height = 高[r]
    for r in range(at, at + n):
        ws.row_dimensions[r].height = None
    for m in 合併:
        if m.min_row >= at:
            ws.merge_cells(start_row=m.min_row + n, start_column=m.min_col, end_row=m.max_row + n, end_column=m.max_col)
        else:
            ws.merge_cells(start_row=m.min_row, start_column=m.min_col, end_row=m.max_row + n, end_column=m.max_col)
    for 表_ in wb.worksheets:
        for row in 表_.iter_rows():
            for 格 in row:
                if isinstance(格.value, str) and 格.value.startswith("="):
                    格.value = 移(格.value, 表_.title)


# 計算表：在合計列下面插入「空一列＋各科目一列＋分類合計＋差異」
#   條件欄：用哪一欄分（月薪表用職稱、獎金表用姓名）；組值：{科目代號: [這個科目的職稱或姓名]}
def 插入計算表(wb, ws, 首, 末, 合計列, 起欄, 迄欄, 條件欄, 組值, 分組):
    標欄 = 起欄 - 1
    插入列(wb, ws, 合計列 + 1, len(分組) + 3)
    上 = 合計列 + 2
    下 = 上 + len(分組) - 1
    條L = get_column_letter(條件欄)
    for i, (標, 代號) in enumerate(分組):
        r = 上 + i
        寫(ws, r, 標欄, 標)
        值們 = 組值.get(代號, [])
        條件 = ",".join('"' + x.replace("~", "~~").replace("*", "~*").replace("?", "~?").replace('"', '""') + '"'
                       for x in 值們)
        for c in range(起欄, 迄欄 + 1):
            L = get_column_letter(c)
            寫(ws, r, c, f"=SUMPRODUCT(SUMIF(${條L}${首}:${條L}${末},{{{條件}}},{L}${首}:{L}${末}))" if 值們 else "=0")
    寫(ws, 下 + 1, 標欄, "分類合計")
    寫(ws, 下 + 2, 標欄, "差異")
    for c in range(起欄, 迄欄 + 1):
        L = get_column_letter(c)
        寫(ws, 下 + 1, c, f"=SUM({L}{上}:{L}{下})")
        寫(ws, 下 + 2, c, f"={L}{合計列}-{L}{下 + 1}")
        for r in range(上, 下 + 3):
            ws.cell(r, c).number_format = ws.cell(合計列, c).number_format


# 科目分類檔：另存一份薪資檔，在每張讀到的表的合計下面插入計算表（原本的資料一格都不改）
def 科目分類檔(src, out, 表們, 對照, 分組, 名科目, 未分類, 警告):
    wb = openpyxl.load_workbook(src)
    做過 = set()
    for t in 表們:
        if t.工作表 in 做過:
            continue
        if not t.合計列:
            警告.append(f"「{t.顯示名稱}」找不到合計列，沒有插入科目計算表")
            continue
        做過.add(t.工作表)
        ws = wb[t.工作表]
        h = t.標題
        首 = t.資料起始列
        末 = max([r for r in range(首, t.合計列) if 是人名(清字(ws.cell(r, t.姓名欄).value))] or [t.合計列 - 1])
        起欄 = min(h[x] for x in t.總額)
        實發 = [c for k, c in h.items() if "實發" in k and c >= 起欄]
        迄欄 = min(實發) if 實發 and t.類型 == "月薪" else max(h.values())
        組值 = OrderedDict()
        用職稱 = t.類型 == "月薪" and "職稱" in h
        條件欄 = h["職稱"] if 用職稱 else t.姓名欄
        for r in range(首, 末 + 1):
            名 = 清字(ws.cell(r, t.姓名欄).value)
            if not 是人名(名):
                continue
            原 = ws.cell(r, 條件欄).value
            科 = 對照.get(清字(原)) if 用職稱 else 名科目.get(名)
            if 科:
                清單 = 組值.setdefault(str(科), [])
                if str(原) not in 清單:
                    清單.append(str(原))
            elif 用職稱:
                未分類.setdefault(清字(原), []).append(f"{t.工作表.strip()} {名}")
            else:
                未分類.setdefault(f"（{t.工作表.strip()} 月薪表找不到職稱）", []).append(名)
        插入計算表(wb, ws, 首, 末, t.合計列, 起欄, 迄欄, 條件欄, 組值, 分組)
    try:
        wb.save(out)
    except PermissionError:
        raise SystemExit(f"無法存檔：{out}\n這個檔案可能正在 Excel 裡開著，請先關掉再執行一次。")
    return sorted(做過, key=lambda x: [t.工作表 for t in 表們].index(x))


# ================================================================ 客戶設定
# 每個客戶一份 .json，放在程式旁邊的「客戶設定」資料夾。第一次執行時，七賢、澤利會自動建好。
空白設定 = {
    "客戶": "", "年度": 115, "薪資檔": "", "樣板": "", "輸出資料夾": "",
    "月薪表": "{年度}年{月}月",
    "薪資總額": "", "退休自提": "自提", "扣繳稅額": "",
    "其他表": [],                         # [工作表名稱, 填入清冊欄位, 月薪／獎金]
    "獎金扣繳詞": "扣繳、國稅局、所得稅、代扣稅",
    "執行業務另開": False,
    "科目分類": False,
    "職稱科目": [],                       # [職稱, 會計科目]
    "計算表分組": "",                     # 「計算表名稱=科目」，例如 6110/6127=6110
    "編號開頭": "",                       # 例如 5300=00 → 0001、0002…
    "科目寫入欄": "扶養人數",              # 清冊上放會計科目的欄
    "人員名冊": "",                       # 選填：清冊的編號、身分證、地址以這份為準
    "模組": "通用",                       # 這個客戶用哪個客戶模組（客戶模組資料夾裡的 .py）
}

def 程式資料夾():
    return os.path.dirname(sys.executable if getattr(sys, "frozen", False) else os.path.abspath(__file__))


def 設定資料夾():
    return os.path.join(程式資料夾(), "客戶設定")


def 安全檔名(名):
    return re.sub(r'[\\/:*?"<>|]', "_", 名.strip()) or "未命名"


# ---------------------------------------------------------------- 客戶模組
# 客戶模組：程式旁邊「客戶模組」資料夾裡的 .py，一個檔案一種讀法（例如 七賢.py、澤利.py、通用.py）。
#   新客戶的格式很特別 → 新增一個 .py 就好，共用核心不用動。模組裡可以有：
#   名稱、說明、預設設定、預設客戶、頁面(介, 設)、整理設定(設)、讀取(設, 年, wb, wbf, 清冊欄)
_模組快取 = OrderedDict()


def 模組資料夾():
    return os.path.join(程式資料夾(), "客戶模組")


def 載入模組們():
    import importlib.util
    sys.modules.setdefault("清冊共用", sys.modules[__name__])   # 讓客戶模組可以 import 清冊共用
    夾 = 模組資料夾()
    if not os.path.isdir(夾):
        return _模組快取
    for f in sorted(os.listdir(夾)):
        if not f.endswith(".py") or f.startswith("_"):
            continue
        鍵 = f[:-3]
        if 鍵 in _模組快取:
            continue
        try:
            規格 = importlib.util.spec_from_file_location(f"客戶模組_{鍵}", os.path.join(夾, f))
            m = importlib.util.module_from_spec(規格)
            _模組快取[鍵] = m              # 先放進去：模組裡可以再叫 取模組("通用") 借用別的模組
            規格.loader.exec_module(m)
        except Exception as e:
            _模組快取.pop(鍵, None)
            print(f"（客戶模組 {f} 載入失敗）", e)
    return _模組快取


def 取模組(鍵):
    return 載入模組們().get(鍵) or 載入模組們().get("通用")


def 新客戶設定(名, 模組鍵):
    m = 取模組(模組鍵)
    設 = copy.deepcopy(空白設定)
    設.update(copy.deepcopy(getattr(m, "預設設定", {})))
    設.update(客戶=名, 模組=模組鍵)
    return 設


def 客戶清單():
    夾 = 設定資料夾()
    os.makedirs(夾, exist_ok=True)
    # 客戶模組的「預設客戶」還沒建過的，自動建好（建過就記在 _已建立.json，刪掉了也不會再冒出來）
    記錄 = os.path.join(夾, "_已建立.json")
    try:
        with open(記錄, encoding="utf-8") as f:
            已建 = set(json.load(f))
    except Exception:
        已建 = set()
    新建 = False
    for 鍵, m in 載入模組們().items():
        for 名 in getattr(m, "預設客戶", []):
            if 名 in 已建:
                continue
            if not os.path.exists(os.path.join(夾, 安全檔名(名) + ".json")):
                存客戶(新客戶設定(名, 鍵))
            已建.add(名)
            新建 = True
    if 新建:
        try:
            with open(記錄, "w", encoding="utf-8") as f:
                json.dump(sorted(已建), f, ensure_ascii=False)
        except Exception:
            pass
    名們 = []
    for f in sorted(os.listdir(夾)):
        if f.endswith(".json") and not f.startswith("_"):
            try:
                with open(os.path.join(夾, f), encoding="utf-8") as fp:
                    名們.append(json.load(fp).get("客戶") or f[:-5])
            except Exception:
                pass
    return 名們


def 讀客戶(名):
    try:
        with open(os.path.join(設定資料夾(), 安全檔名(名) + ".json"), encoding="utf-8") as f:
            存的 = json.load(f)
    except FileNotFoundError:
        存的 = {}
    設 = 新客戶設定(名, 存的.get("模組", "通用"))
    設.update(存的)
    設["客戶"] = 名
    return 設


def 存客戶(設):
    os.makedirs(設定資料夾(), exist_ok=True)
    try:
        with open(os.path.join(設定資料夾(), 安全檔名(設["客戶"]) + ".json"), "w", encoding="utf-8") as f:
            json.dump(設, f, ensure_ascii=False, indent=1)
    except Exception as e:
        print("（客戶設定存檔失敗）", e)


def 刪客戶(名):
    p = os.path.join(設定資料夾(), 安全檔名(名) + ".json")
    if os.path.exists(p):
        os.remove(p)


def 上次客戶(新=None):
    p = os.path.join(程式資料夾(), "清冊共用_上次客戶.json")
    if 新 is not None:
        try:
            with open(p, "w", encoding="utf-8") as f:
                json.dump({"上次客戶": 新}, f, ensure_ascii=False)
        except Exception:
            pass
        return 新
    try:
        with open(p, encoding="utf-8") as f:
            return json.load(f).get("上次客戶", "")
    except Exception:
        return ""


# ================================================================ 人員名冊（選用）
# 開活頁簿：.xlsx 用 openpyxl 開；舊版 .xls 用 xlrd 讀，再轉成 openpyxl 的格式（只要值）
# 「big5寬鬆」編碼：用 Big5-HKSCS 解，解不開的字換成「□」（給讀舊版 .xls 用）
def _big5寬鬆(名稱):
    if 名稱 != "x_big5_loose":
        return None
    import codecs
    底 = codecs.lookup("big5hkscs")

    def 解(資料, errors="strict"):
        文字, 長 = 底.decode(資料, "replace")
        return 文字.replace("\ufffd", "□"), 長
    return codecs.CodecInfo(底.encode, 解, name="x_big5_loose")


import codecs as _codecs
_codecs.register(_big5寬鬆)


def 開活頁簿(路徑):
    if not 路徑.lower().endswith(".xls"):
        return openpyxl.load_workbook(路徑, data_only=True)
    try:
        import xlrd
    except ImportError:
        raise SystemExit("這是舊版的 .xls 檔，需要先安裝 xlrd：pip install xlrd\n（或用 Excel 把它另存成 .xlsx）")
    書 = None
    # 舊版 .xls 的中文是 Big5（cp950）。有些檔案有造字、罕用字（例如人名用字），cp950 解不開，
    #   就依序改用 Big5-HKSCS（字比較多）→ 還是不行的話，解不開的字先用「□」代替，至少讀得到其他資料
    for 編碼 in (None, "big5hkscs", "x_big5_loose"):
        try:
            書 = xlrd.open_workbook(路徑, encoding_override=編碼)
            break
        except UnicodeDecodeError:
            continue
        except Exception as e:
            raise SystemExit(f".xls 讀不到：{e}\n可以用 Excel 開啟後另存成 .xlsx 再試一次")
    if 書 is None:
        raise SystemExit(".xls 的中文編碼讀不到，請用 Excel 開啟後另存成 .xlsx 再試一次")
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for 表_ in 書.sheets():
        ws = wb.create_sheet(表_.name[:31])
        for r in range(表_.nrows):
            for c in range(表_.ncols):
                v = 表_.cell_value(r, c)
                if v == "":
                    continue
                if isinstance(v, float) and v.is_integer():   # xlrd 的數字都是小數：1.0 → 1
                    v = int(v)
                ws.cell(r + 1, c + 1).value = v
    return wb


# 讀前年度：去年的印領清冊、名冊都可以 —— 任何工作表只要前 10 列有「姓名」標題就讀，
#   抓 姓名、身分證號、員工編號、地址。回傳 {姓名: {身分證號, 編號, 地址, 來源}}
def 讀前年度(路徑):
    wb = 開活頁簿(路徑)
    人 = OrderedDict()
    for ws in wb.worksheets:
        if ws.title.strip() in ("匯入檢查", "前年度核對", "人員名冊核對"):
            continue
        for r in range(1, 11):
            h = {清字(ws.cell(r, c).value): c for c in range(1, min(ws.max_column, 60) + 1) if 清字(ws.cell(r, c).value)}
            名 = next((h[k] for k in 姓名標題 if k in h), None)
            if not 名:
                continue
            證 = next((c for k, c in h.items() if any(x in k for x in 身分證關鍵字)), None)
            編 = next((h[k] for k in 編號標題 if k in h), None)
            址 = next((c for k, c in h.items() if "地址" in k), None)
            類代 = next((c for k, c in h.items() if "類別" in k and "代號" in k), None)
            類 = next((c for k, c in h.items() if "類別" in k and "代號" not in k), None)
            for rr in range(r + 1, ws.max_row + 1):
                n = 清字(ws.cell(rr, 名).value)
                if not 是人名(n):
                    continue
                新 = {"身分證號": 清字(ws.cell(rr, 證).value).upper() if 證 else "",
                     "編號": ws.cell(rr, 編).value if 編 else None,
                     "地址": 清字(ws.cell(rr, 址).value) if 址 else "", "來源": ws.title.strip(), "類別": [],
                     "地址原": str(ws.cell(rr, 址).value).strip() if 址 and ws.cell(rr, 址).value is not None else ""}
                類字 = "／".join(x for x in (清字(ws.cell(rr, 類代).value) if 類代 else "",
                                           清字(ws.cell(rr, 類).value) if 類 else "") if x)
                舊 = 人.setdefault(n, 新)
                if 類字 and 類字 not in 舊["類別"]:          # 同一人可能有兩種類別（薪資＋執行業務）
                    舊["類別"].append(類字)
                if not 類 and not 類代:                   # 沒有類別欄的表（例如去年的清冊）：用工作表名稱判斷
                    類字 = "執行業務" if "執行業務" in ws.title else "薪資"
                    if 類字 not in 舊["類別"]:
                        舊["類別"].append(類字)
                for k in ("身分證號", "編號", "地址", "地址原"):     # 同一人出現在好幾張表：空的欄位用後面的補
                    if not 舊.get(k) and 新.get(k):
                        舊[k] = 新[k]
            break
    if not 人:
        raise SystemExit("前年度資料裡找不到有「姓名」標題的工作表（例如去年的印領清冊或薪資名冊）")
    return 人


# 地址鍵：比地址前先統一寫法 —— 全形數字／英文轉半形（１０鄰 → 10鄰）、臺 → 台、去掉空白和符號差異
def 地址鍵(x):
    import unicodedata
    t = unicodedata.normalize("NFKC", 清字(x))
    t = t.replace("臺", "台").replace("－", "-").replace("—", "-")
    return re.sub(r"[\s,，、.。]", "", t).upper()


# 類別歸類：名冊上的所得類別 → 薪資／執行業務／其他（空白回傳 None，不比）
#   例：「50-1／薪資」→ 薪資；「9A42／執行業務(醫事檢驗師)」→ 執行業務；
#       「928Z／其他所得」「97／受贈所得」「91-1／競技競賽中獎」→ 其他（不是印領清冊的範圍）
def 類別歸類(x):
    t = 清字(x).upper()
    if not t:
        return None
    if "執行" in t or re.match(r"^9[AB]", t):
        return "執行業務"
    if "薪" in t or re.match(r"^50(\D|$)", t):
        return "薪資"
    return "其他"


# 前年度核對：今年清冊上的人 vs 前年度資料，回傳 OrderedDict{類別: [(姓名, 今年, 去年, 說明)]}
#   先用身分證號對，對不到再用姓名對
def 前年度核對(今年, 前年, 比地址=True):
    結果 = OrderedDict((k, []) for k in ("同身分證、名字不同", "同名字、身分證不同", "客戶資料有缺",
                                        "所得類別不同", "地址不同", "名冊找不到", "名冊有、清冊沒有",
                                        "名冊上是其他所得"))
    依證 = {v["身分證號"]: n for n, v in 前年.items() if v.get("身分證號")}
    對到, 新人 = set(), []
    for n, info, 今類 in 今年:
        證 = 清字(info.get("身分證號")).upper()
        舊名 = 依證.get(證) if 證 else None
        if 舊名 is None and n in 前年:
            舊名 = n
        if 舊名 is None:
            新人.append((n, 證))
            continue
        對到.add(舊名)
        舊 = 前年[舊名]
        if 舊名 != n:
            結果["同身分證、名字不同"].append((n, n, 舊名, "可能改名或打錯字"))
        舊證 = 舊.get("身分證號", "")
        if 證 and 舊證 and 證 != 舊證:
            結果["同名字、身分證不同"].append((n, 證, 舊證, "清冊用人員名冊的；客戶薪資名冊可能打錯，或同名不同人"))
        elif not 證 and 舊證:
            結果["客戶資料有缺"].append((n, "", 舊證, "客戶沒有身分證號，清冊用人員名冊的"))
        if 比地址 and not 清字(info.get("地址")) and 舊.get("地址"):
            結果["客戶資料有缺"].append((n, "", 舊.get("地址原") or 舊.get("地址"), "客戶沒有地址，清冊用人員名冊的"))
        舊類 = 舊.get("類別") or []
        舊類別 = {類別歸類(x) for x in 舊類} - {None}
        if 舊類別 and 今類 not in 舊類別:
            結果["所得類別不同"].append((n, 今類, "、".join(舊類), f"今年放在「{今類}」，名冊上是「{'、'.join(舊類)}」"
                                       + ("（名冊上不是薪資或執行業務，請確認要不要放進清冊）" if 舊類別 == {"其他"} else "")))
        # 編號不比：今年是客戶名冊的員工編號，對照檔的編號通常是另一套（例如申報檔的流水號）
        變 = []
        if 清字(info.get("地址")) and 舊.get("地址") and 地址鍵(info.get("地址")) != 地址鍵(舊.get("地址")):
            變.append(("地址", info.get("地址"), 舊.get("地址")))
        for 項, a, b in 變:
            結果["地址不同"].append((n, a, b, "清冊用人員名冊的"))
    沒對到 = [x for x in 前年 if x not in 對到]      # 新人：只跟「去年有、今年沒對到」的人比相近名字
    for n, 證 in 新人:
        像 = difflib.get_close_matches(n, 沒對到, n=2, cutoff=0.6)
        結果["名冊找不到"].append((n, 證, "", ("⚠ 去年有相近名字：" + "、".join(像) + "，請確認是不是同一人") if 像 else ""))
    for n, v in 前年.items():
        if n in 對到:
            continue
        類們 = v.get("類別") or []
        if 類們 and {類別歸類(x) for x in 類們} == {"其他"}:     # 其他所得、受贈、中獎…本來就不在印領清冊
            結果["名冊上是其他所得"].append((n, "", "、".join(類們), "不是薪資或執行業務，清冊沒有是正常的"))
        else:
            結果["名冊有、清冊沒有"].append((n, "", v.get("身分證號", ""),
                                       "可能離職，或今年名字寫法不同" + (f"（{'、'.join(類們)}）" if 類們 else "")))
    return 結果


# 寫人員名冊核對工作表：上面一張「摘要」（每類幾人），下面一張明細表（可以用篩選只看某一類）
def 寫前年度核對(wbo, 結果, 路徑, 人數):
    from openpyxl.styles import Alignment, Border, Side
    ws = wbo.create_sheet("人員名冊核對")
    ws.sheet_view.showGridLines = False
    說明 = OrderedDict([
        ("同身分證、名字不同", "身分證號一樣、名字不一樣：可能改名或打錯字"),
        ("同名字、身分證不同", "名字一樣、身分證號不一樣：客戶可能打錯，或同名不同人"),
        ("所得類別不同", "名冊上的所得類別跟清冊放的工作表（薪資／執行業務）不一樣"),
        ("地址不同", "兩邊都有地址但不一樣（全形／半形數字、臺／台視為一樣）"),
        ("客戶資料有缺", "客戶沒有身分證號或地址，名冊有"),
        ("名冊找不到", "人員名冊沒有這個人：清冊上他的基本資料用客戶的，請確認"),
        ("名冊有、清冊沒有", "名冊上是薪資或執行業務，但今年清冊沒有：可能離職或名字寫法不同"),
        ("名冊上是其他所得", "其他所得、受贈、中獎這類，本來就不在印領清冊"),
    ])
    重要 = {"同身分證、名字不同", "同名字、身分證不同", "所得類別不同", "名冊找不到"}
    奶茶, 淡奶茶, 橘, 灰字 = "D7B89F", "F6EEE4", "FCE4D0", "A69486"
    細 = Side(style="thin", color="E2D3C3")
    框 = Border(left=細, right=細, top=細, bottom=細)
    置中 = Alignment(horizontal="center", vertical="center")
    換行 = Alignment(vertical="center", wrap_text=True)
    for col, w in zip("ABCDE", (20, 11, 34, 34, 46)):
        ws.column_dimensions[col].width = w

    # 標題
    ws["A1"] = "人員名冊核對"
    ws["A1"].font = Font(size=15, bold=True, color="5A4636")
    ws["A2"] = (f"人員名冊：{os.path.basename(路徑)}（{人數} 人）｜清冊的編號、身分證號、地址以人員名冊為準，"
                f"「客戶資料」是客戶薪資名冊上寫的")
    ws["A2"].font = Font(size=10, color=灰字)

    # 摘要
    r = 4
    for c, t in enumerate(("項目", "人數", "說明"), 1):
        格 = ws.cell(r, c if c < 3 else 3, t)
        格.font = Font(bold=True, color="FFFFFF"); 格.fill = PatternFill("solid", fgColor=奶茶); 格.alignment = 置中
    ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
    for 類, 說 in 說明.items():
        r += 1
        n = len(結果.get(類, []))
        ws.cell(r, 1, 類); ws.cell(r, 2, n); ws.cell(r, 3, 說)
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=5)
        for c in range(1, 6):
            ws.cell(r, c).border = 框
            ws.cell(r, c).alignment = 置中 if c == 2 else 換行
        if n and 類 in 重要:
            for c in range(1, 6):
                ws.cell(r, c).fill = PatternFill("solid", fgColor=橘)
            ws.cell(r, 1).font = Font(bold=True)
        elif not n:
            for c in range(1, 6):
                ws.cell(r, c).font = Font(color=灰字)

    # 明細
    r += 2
    表頭 = r
    for c, t in enumerate(("項目", "姓名", "客戶資料", "人員名冊", "說明"), 1):
        格 = ws.cell(r, c, t)
        格.font = Font(bold=True, color="FFFFFF"); 格.fill = PatternFill("solid", fgColor=奶茶)
        格.alignment = 置中; 格.border = 框
    for 類 in 說明:
        for i, 列 in enumerate(結果.get(類, [])):
            r += 1
            n, 今, 舊, 說 = 列
            for c, v in enumerate((類, n, 今, 舊, 說), 1):
                格 = ws.cell(r, c, "" if v is None else v)
                格.border = 框
                格.alignment = 換行
            if 類 in 重要 or str(說).startswith("⚠"):
                for c in range(1, 6):
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=橘)
            elif i % 2:
                for c in range(1, 6):
                    ws.cell(r, c).fill = PatternFill("solid", fgColor=淡奶茶)
    if r > 表頭:
        ws.auto_filter.ref = f"A{表頭}:E{r}"          # 表頭有篩選按鈕，可以只看某一類
    ws.freeze_panes = f"A{表頭 + 1}"
    ws.row_dimensions[1].height = 24
    ws.page_setup.orientation = "landscape"           # 列印：橫向、寬度縮成一頁
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = f"{表頭}:{表頭}"


# ================================================================ 給客戶模組用：自己讀的資料
# 資料來源：客戶模組自己讀檔案（例如每月一個 Excel、PDF、一張表放全年）時，把讀到的每個人放進「列們」，
#   共用核心就會接手：歸戶、填清冊、驗算、匯入檢查、人員名冊核對
#   類型：「月薪」（填一月～十二月）或「獎金」；欄位：清冊上的欄名（例如 三月、上期年終、端午獎金）
class 資料來源:
    def __init__(self, 名稱, 欄位, 類型="月薪"):
        self.工作表 = 名稱
        self.顯示名稱 = 名稱
        self.欄位 = 清字(欄位)
        self.類型 = 類型
        self.總額, self.自提, self.扣繳 = [], None, []
        self.核對 = []
        self.合計列 = None
        self.列們 = []

    # 加一個人：名＝姓名；總額、自提、稅＝金額；證＝身分證；職＝職稱；編＝員工編號；執＝是不是執行業務
    def 加(self, 名, 總額=0, 自提=0, 稅=0, 證="", 職="", 編="", 執=False):
        self.列們.append({"名": 清字(名), "編": 編號鍵(編), "金": {"總額": 數字(總額), "自提": 數字(自提), "稅": 數字(稅)},
                        "執": bool(執), "職": 清字(職), "證": 清字(證).upper()})

    # 核對：讀完後跟表上的合計比，對不起來會寫進匯入檢查
    def 對合計(self, 項目, 表上合計):
        鍵 = {"總額": "總額", "自提": "自提", "稅": "稅"}[項目]
        self.核對.append((項目, sum(x["金"][鍵] for x in self.列們), 數字(表上合計)))


# ================================================================ 主流程
def 準備樣板(tpl, 新增欄, 要稅, 月份們=None):
    wbo = openpyxl.load_workbook(tpl)
    ws = 樣板工作表(wbo)
    表頭列, _, 清冊欄 = 清冊欄位表(ws)
    if 新增欄:
        位置 = 清冊欄.get("本期年終") or 清冊欄.get("合計") or (max(清冊欄.values()) + 1)
        for 名 in 新增欄:
            插入欄(ws, 位置, 名, 表頭列)
            位置 += 1
    樣 = 樣板(ws, 要稅)
    樣.加獎金列(list(月份們 or {}))
    return wbo, ws, 樣


def 執行(設, 輸出資料夾, 模組=None):
    模組 = 模組 or 取模組(設.get("模組", "通用"))
    if 模組 is not None and hasattr(模組, "整理設定"):      # 客戶模組把自己的設定換成共用的讀法
        設 = 模組.整理設定(copy.deepcopy(設))
    自己讀 = 模組 is not None and hasattr(模組, "讀取")
    globals()["月薪金額詞"] = 拆詞(設.get("薪資總額", ""))[:1] or ["扣憑"]
    年 = int(設["年度"])
    src, tpl = 設["薪資檔"], 設["樣板"]
    if not src or not os.path.exists(src):
        raise SystemExit("請先選擇薪資檔")
    if not tpl or not os.path.exists(tpl):
        raise SystemExit("請先選擇清冊樣板")
    if not 自己讀 and not 拆詞(設.get("薪資總額", "")):
        raise SystemExit("請在「月薪表」填上薪資總額是哪幾欄（例如 應付薪資＋員工酬勞）")
    os.makedirs(輸出資料夾, exist_ok=True)
    客 = 安全檔名(設["客戶"])

    # ---- 1. 樣板：設定裡有、樣板沒有的欄位 → 之後自動插在「本期年終」前面
    wbo0 = openpyxl.load_workbook(tpl)
    _, _, 清冊欄 = 清冊欄位表(樣板工作表(wbo0))
    新增欄 = []
    for x in 設.get("其他表", []):
        if 清字(x[1]) and 清字(x[1]) not in 清冊欄 and 清字(x[1]) not in 新增欄:
            新增欄.append(清字(x[1]))
    清冊欄 = OrderedDict(清冊欄)
    for x in 新增欄:
        清冊欄[x] = 0
    # 選了月份的獎金（其他表第 4 格，例如春節獎金放 2 月）：每人多一列，值放在那個月
    月份們 = OrderedDict()
    for x in 設.get("其他表", []):
        if len(x) > 3 and x[3]:
            月份們[清字(x[1])] = int(str(x[3]).replace("月", ""))

    # ---- 2. 讀薪資檔
    if src.lower().endswith(".xls"):          # 舊版 .xls：只讀得到值（沒有公式）
        wb = wbf = 開活頁簿(src)
    else:
        wb = openpyxl.load_workbook(src, data_only=True)
        wbf = openpyxl.load_workbook(src)
    名冊表, 名冊, 編號對照 = 讀名冊(wb)
    for 名 in 月份們:
        if 名 not in 清冊欄:
            raise SystemExit(f"獎金欄位「{名}」不在樣板裡")
    if 自己讀:
        表們, 缺表, 警告 = 模組.讀取(設, 年, wb, wbf, 清冊欄)
    else:
        表們, 缺表, 警告 = 選表(wb, 設, 年, 清冊欄)
    if not any(t.類型 == "月薪" for t in 表們):
        raise SystemExit(f"薪資檔裡找不到 {年} 年的月薪表（工作表名稱像「{表名(設['月薪表'], 年, 1)}」），"
                         f"請檢查客戶設定的「月薪表名稱」")
    原始 = []
    for t in 表們:
        for x in (t.列們 if getattr(t, "列們", None) is not None else 讀表資料(t, wbf, 警告)):
            原始.append((t, x))

    # ---- 3. 人員歸戶：有員工編號而且名冊找得到 → 用編號；否則用姓名
    改名 = {}
    for t, x in 原始:
        if x["編"] in 編號對照 and 編號對照[x["編"]] != x["名"]:
            改名.setdefault(x["名"], (編號對照[x["編"]], x["編"], t.顯示名稱))

    def 歸戶(n, 編):
        if 編 in 編號對照:
            return 編號對照[編]
        return 改名[n][0] if n in 改名 else n
    用名 = {}
    for t, x in 原始:
        用名.setdefault(歸戶(x["名"], x["編"]), Counter())[x["名"]] += 1
    顯示名 = {k: c.most_common(1)[0][0] for k, c in 用名.items()}

    # ---- 4. 基本資料：名冊優先；名冊沒有身分證的，用薪資表上的身分證欄（越後面的月份優先）
    基本 = {n: dict(v) for n, v in 名冊.items()}
    表上證 = {}
    for t, x in 原始:
        if x["證"] and len(x["證"]) >= 8:
            表上證[歸戶(x["名"], x["編"])] = x["證"]
    備用證 = None
    for 人 in 用名:
        b = 基本.setdefault(人, {})
        if not b.get("身分證號"):
            if 人 in 表上證:
                b["身分證號"] = 表上證[人]
            else:
                if 備用證 is None:
                    備用證 = 讀身分證(wb)
                if 人 in 備用證:
                    b["身分證號"] = 備用證[人]

    # ---- 5. 分到薪資／執行業務，並累計核對數字
    執行業務另開 = bool(設.get("執行業務另開"))
    有薪, 有執 = set(), set()
    for t, x in 原始:
        if t.類型 == "月薪":
            (有執 if (x["執"] and 執行業務另開) else 有薪).add(歸戶(x["名"], x["編"]))
    純執 = (有執 - 有薪) if 執行業務另開 else set()
    薪資, 執行 = OrderedDict(), OrderedDict()
    核對 = OrderedDict()
    獎金轉執行 = []
    for t, x in 原始:
        人 = 歸戶(x["名"], x["編"])
        去執 = 執行業務另開 and (x["執"] or (t.類型 != "月薪" and 人 in 純執))
        目的 = 執行 if 去執 else 薪資
        d = 目的.setdefault(人, OrderedDict()).setdefault(t.欄位, {"總額": 0, "自提": 0, "稅": 0})
        for k in d:
            d[k] += x["金"][k]
        c = 核對.setdefault((t, "執行業務" if 去執 else "薪資"), {"人": set(), "總額": 0, "自提": 0, "稅": 0})
        c["人"].add(人)
        for k in ("總額", "自提", "稅"):
            c[k] += x["金"][k]
        if 去執 and t.類型 != "月薪":
            獎金轉執行.append(f"{t.顯示名稱}：{顯示名.get(人, 人)} {x['金']['總額']:,}"
                           + (f"（扣繳 {x['金']['稅']:,}）" if x["金"]["稅"] else "")
                           + ("，表上寫明執行業務" if x["執"] else "，月薪表只有執行業務"))

    def 名單(資料):
        有 = [n for n, d in 資料.items() if any(v["總額"] or v["自提"] or v["稅"] for v in d.values())]
        單 = [n for n in 名冊 if n in 有 or (not 只列有資料的人 and 資料 is 薪資)]
        外 = [n for n in 有 if n not in 名冊]
        return 單 + 外, 外
    薪單, 薪外 = 名單(薪資)
    執單, 執外 = 名單(執行) if 執行 else ([], [])
    要稅 = any(v["稅"] for d in 薪資.values() for v in d.values()) or bool(拆詞(設.get("扣繳稅額", "")))

    # ---- 6. 科目分類（選用）：職稱 → 會計科目；另存科目分類檔；清冊依科目分檔
    未分類, 科目of, 分類檔, 分類表 = OrderedDict(), {}, None, []
    if 設.get("科目分類"):
        對照 = {清字(a): 清字(b) for a, b in 設.get("職稱科目", []) if 清字(a) and 清字(b)}
        for t, x in 原始:                      # 用月薪表的職稱（最後出現的月份為準）
            if t.類型 == "月薪" and x["職"] and 對照.get(x["職"]):
                科目of[歸戶(x["名"], x["編"])] = 對照[x["職"]]
        分組 = 拆對照(設.get("計算表分組", "")) or [(k, k) for k in dict.fromkeys(對照.values())]
        基名, 副 = os.path.splitext(os.path.basename(src))
        分類檔 = os.path.join(輸出資料夾, f"{基名}_科目分類{副 or '.xlsx'}")
        if os.path.abspath(分類檔) == os.path.abspath(src):
            raise SystemExit("輸出資料夾不能讓科目分類檔蓋掉原本的薪資檔")
        名科目 = {}                             # 表上的每一種寫法 → 會計科目（獎金表依姓名分科目用）
        for k, c in 用名.items():
            for 寫法 in c:
                if k in 科目of:
                    名科目[寫法] = 科目of[k]
        try:
            分類表 = 科目分類檔(src, 分類檔, 表們, 對照, 分組, 名科目, 未分類, 警告)
        except SystemExit:
            raise
        except Exception as e:
            raise SystemExit(f"在「科目分類」這一步出錯：\n{type(e).__name__}: {e}")

    # ---- 6.5 人員名冊（選用）：清冊的編號、身分證號、地址改用名冊的；先用客戶的資料跟名冊核對
    前年, 核, 用客戶 = None, None, []
    if 設.get("人員名冊"):
        if not os.path.exists(設["人員名冊"]):
            raise SystemExit(f"找不到人員名冊：{設['人員名冊']}")
        前年 = 讀前年度(設["人員名冊"])
        if 設.get("科目分類"):                # 分科目時，沒有科目的人不會放進清冊，也不用核對
            薪放 = [n for n in 薪單 if re.match(r"\d+", str(科目of.get(n, "")))]
        else:
            薪放 = 薪單
        今年 = ([(顯示名.get(n, n), 基本.get(n, {}), "薪資") for n in 薪放]
              + [(顯示名.get(n, n), 基本.get(n, {}), "執行業務") for n in 執單])
        核 = 前年度核對(今年, 前年, 比地址=bool(名冊))
        依證 = {v["身分證號"]: k for k, v in 前年.items() if v.get("身分證號")}
        for n in 薪放 + 執單:
            客資 = 基本.get(n, {})
            m = next((x for x in (顯示名.get(n, n), n) if x in 前年), None) or 依證.get(清字(客資.get("身分證號")).upper())
            if m:
                v = 前年[m]
                基本[n] = {"編號": v.get("編號"), "身分證號": v.get("身分證號") or None, "地址": v.get("地址原") or None}
            else:
                用客戶.append(顯示名.get(n, n))

    # ---- 7. 填清冊
    輸出們 = []
    if 設.get("科目分類"):
        開頭 = dict(拆對照(設.get("編號開頭", "")))
        組們 = OrderedDict()
        沒科目 = []
        for n in 薪單:
            科 = re.match(r"\d*", str(科目of.get(n, ""))).group(0)
            (組們.setdefault(科, []) if 科 else 沒科目).append(n)
        for 科 in sorted(組們):
            wbo, ws, 樣 = 準備樣板(tpl, 新增欄, 要稅, 月份們)
            編號們 = {n: f"{開頭[科]}{i + 1:02d}" for i, n in enumerate(組們[科])} if 科 in 開頭 else {}
            樣.填(組們[科], 薪資, 基本, 顯示名, 編號們, {n: 科 for n in 組們[科]}, 設.get("科目寫入欄", ""), 月份們)
            輸出們.append((os.path.join(輸出資料夾, f"{客}_{年}年薪資印領清冊_{科}.xlsx"), wbo, f"科目 {科}：{len(組們[科])} 人"))
        if 沒科目:
            警告.append("這些人的職稱沒有對應會計科目，沒有放進清冊：" + "、".join(顯示名.get(n, n) for n in 沒科目))
        if 執單:
            wbo, ws, 樣 = 準備樣板(tpl, 新增欄, 要稅, 月份們)
            樣.填(執單, 執行, 基本, 顯示名, 月份們=月份們)
            輸出們.append((os.path.join(輸出資料夾, f"{客}_{年}年執行業務印領清冊.xlsx"), wbo, f"執行業務：{len(執單)} 人"))
    else:
        wbo, ws, 樣 = 準備樣板(tpl, 新增欄, 要稅, 月份們)
        ws執 = None
        if 執單:
            ws執 = wbo.copy_worksheet(ws)
            ws執.title = "執行業務"
            if isinstance(ws執["A1"].value, str):
                ws執["A1"].value = ws執["A1"].value.replace("薪資印領清冊", "執行業務所得印領清冊")
        樣.填(薪單, 薪資, 基本, 顯示名, 月份們=月份們)
        if ws執 is not None:
            樣執 = 樣板.__new__(樣板)
            樣執.__dict__.update(樣.__dict__)
            樣執.ws = ws執
            樣執.填(執單, 執行, 基本, 顯示名, 月份們=月份們)
        輸出們.append((os.path.join(輸出資料夾, f"{客}_{年}年薪資印領清冊.xlsx"), wbo,
                     f"薪資 {len(薪單)} 人" + (f"，執行業務 {len(執單)} 人" if 執單 else "")))
    for p, _, _ in 輸出們:
        if os.path.abspath(p) in (os.path.abspath(tpl), os.path.abspath(src)):
            raise SystemExit("輸出檔不能和薪資檔或樣板同一個檔案")

    # ---- 8. 匯入檢查
    放進的人 = 薪單 + 執單 if not 設.get("科目分類") else [n for n in 薪單 if n in 科目of] + 執單
    缺證 = [顯示名.get(n, n) for n in 放進的人 if not 基本.get(n, {}).get("身分證號")]
    只在獎金 = []
    if not 名冊:
        月薪人 = {歸戶(x["名"], x["編"]) for t, x in 原始 if t.類型 == "月薪"}
        for n in 用名:
            if n not in 月薪人:
                像 = difflib.get_close_matches(n, list(月薪人), n=2, cutoff=0.5)
                只在獎金.append((顯示名.get(n, n), 像))

    def 寫檢查(wbo):
        chk = wbo.create_sheet("匯入檢查")
        for col, w in zip("ABCDEFG", (30, 12, 10, 14, 12, 12, 80)):
            chk.column_dimensions[col].width = w
        粗, 橘 = Font(bold=True), PatternFill("solid", fgColor="FFF3E0")

        def 標題(*v):
            chk.append(list(v))
            for c in chk[chk.max_row]:
                c.font = 粗

        def 橘列():
            for c in chk[chk.max_row]:
                c.fill = 橘
        標題(f"客戶：{設['客戶']}　清冊年度：{年}")
        chk.append([])
        標題("【對應與核對】", "填入欄位", "人數", "總額", "自提", "扣繳", "工作表／欄位／表上合計核對")
        for t in 表們:
            for 類 in ("薪資", "執行業務"):
                c = 核對.get((t, 類))
                if not c:
                    continue
                錯 = [f"{標} {a:,}≠{b:,}" for 標, a, b in getattr(t, "核對", []) if abs(a - b) > 1]
                說明 = (f"總額：{'＋'.join(t.總額)}" + (f"；自提：{t.自提}" if t.自提 else "")
                      + f"；扣繳：{'＋'.join(t.扣繳) or '無'}"
                      + ("" if not getattr(t, "核對", None) or 類 != "薪資" else
                         ("；合計核對 ✓" if not 錯 else "；⚠ 合計不符：" + "、".join(錯))))
                chk.append([t.顯示名稱 + ("（執行業務）" if 類 == "執行業務" else ""), t.欄位, len(c["人"]),
                            c["總額"], c["自提"], c["稅"], 說明])
                if 錯 and 類 == "薪資":
                    橘列()
        for n in 缺表:
            chk.append([n, None, None, None, None, None, "⚠ 薪資檔裡找不到這張工作表，略過"])
            橘列()
        for row in chk.iter_rows(min_row=3, min_col=4, max_col=6):
            for c in row:
                c.number_format = "#,##0"
        if 名冊 and (薪外 or 執外):
            chk.append([])
            標題("【不在名冊】", f"{len(薪外) + len(執外)} 人", None, None, None, None, "編號、地址空白，請補到名冊（有相近名字的請確認是不是打錯字）")
            for 類, 外, 資料 in (("薪資", 薪外, 薪資), ("執行業務", 執外, 執行)):
                for n in 外:
                    像 = difflib.get_close_matches(n, list(名冊), n=3, cutoff=0.6)
                    chk.append([顯示名.get(n, n), 類, None, None, None, None,
                                ("⚠ 名冊有相近名字：" + "、".join(像) + "；" if 像 else "") + "出現在：" + "、".join(資料[n].keys())])
                    if 像:
                        橘列()
        if 只在獎金:
            chk.append([])
            標題("【只出現在獎金表】", f"{len(只在獎金)} 人", None, None, None, None, "月薪表沒有這個人，請確認是不是名字打錯")
            for n, 像 in 只在獎金:
                chk.append([n, None, None, None, None, None, ("⚠ 月薪表有相近名字：" + "、".join(像)) if 像 else ""])
                if 像:
                    橘列()
        if 缺證:
            chk.append([])
            標題("【沒有身分證號】", f"{len(缺證)} 人", None, None, None, None, "清冊的身分證號空白，請補到名冊或薪資表")
            for i in range(0, len(缺證), 10):
                chk.append(["", None, None, None, None, None, "、".join(缺證[i:i + 10])])
        if 改名:
            chk.append([])
            標題("【編號相同、姓名不同】", f"{len(改名)} 筆", None, None, None, None, "已依編號視為同一人")
            for n, (名冊名, 編, 表名_) in 改名.items():
                chk.append([n, None, None, None, None, None, f"編號 {編}：名冊為「{名冊名}」（見 {表名_}）"])
        if 獎金轉執行:
            chk.append([])
            標題("【獎金放到執行業務】", f"{len(獎金轉執行)} 筆")
            for x in 獎金轉執行:
                chk.append(["", None, None, None, None, None, x])
        if 未分類:
            chk.append([])
            標題("【沒有對應會計科目】", None, None, None, None, None, "沒算進科目計算表（差異不會是 0），也不會放進清冊")
            for k, v in 未分類.items():
                chk.append([k, f"{len(v)} 筆", None, None, None, None, "、".join(v[:8]) + ("…" if len(v) > 8 else "")])
                橘列()
        if 警告:
            chk.append([])
            標題("【請確認】")
            for x in 警告:
                chk.append(["", None, None, None, None, None, x])

    for p, wbo, _ in 輸出們:
        寫檢查(wbo)
        if 前年 is not None:
            寫前年度核對(wbo, 核, 設["人員名冊"], len(前年))
        try:
            wbo.save(p)
        except PermissionError:
            raise SystemExit(f"無法存檔：{p}\n這個檔案可能正在 Excel 裡開著，請先關掉再執行一次。")

    # ---- 9. 摘要
    月份 = [t.工作表.strip() for t in 表們 if t.類型 == "月薪" and t.欄位 in 中文月]
    摘要 = [f"輸出資料夾：{輸出資料夾}"]
    if 分類檔:
        摘要.append(f"・{os.path.basename(分類檔)}：{len(分類表)} 張表插入科目計算表")
    摘要 += [f"・{os.path.basename(p)}：{說}" for p, _, 說 in 輸出們]
    摘要.append("")
    摘要.append(f"月薪表：{月份[0]}～{月份[-1]} 共 {len(月份)} 張" if 月份 else "月薪表：0 張")
    不符 = [t.顯示名稱 for t in 表們 if any(abs(a - b) > 1 for _, a, b in getattr(t, "核對", []))]
    有核 = [t for t in 表們 if getattr(t, "核對", None)]
    if 有核:
        摘要.append("每張表的個人加總跟表上合計：" + ("全部一致 ✓" if not 不符 else f"⚠ {len(不符)} 張不一致：" + "、".join(不符)))
    if 新增欄:
        摘要.append(f"已在清冊新增欄位：{'、'.join(新增欄)}")
    if 缺表:
        摘要.append(f"⚠ 找不到工作表：{'、'.join(缺表)}")
    if 名冊 and (薪外 or 執外):
        摘要.append(f"⚠ {len(薪外) + len(執外)} 人不在名冊")
    if 只在獎金:
        摘要.append(f"⚠ {len(只在獎金)} 人只出現在獎金表：" + "、".join(n for n, _ in 只在獎金))
    if 缺證:
        摘要.append(f"⚠ {len(缺證)} 人沒有身分證號")
    if 改名:
        摘要.append(f"⚠ {len(改名)} 人編號相同但姓名不同（已合併）")
    if 未分類:
        摘要.append("⚠ 沒有對應會計科目：" + "、".join(k for k in 未分類))
    摘要 += [f"⚠ {x}" for x in 警告 if "不符" not in x and "≠" not in x]
    if 前年 is not None:
        摘要 += ["", f"人員名冊：清冊的編號、身分證、地址 {len(薪放) + len(執單) - len(用客戶)} 人用名冊"
                 + (f"，{len(用客戶)} 人名冊找不到（用客戶的資料）" if 用客戶 else ""),
                 "人員名冊核對：" + ("、".join(f"{k} {len(v)} 人" for k, v in 核.items() if v) or "全部一致 ✓")]
    摘要.append("\n詳細內容請看輸出檔的「匯入檢查」" + ("、「人員名冊核對」" if 前年 is not None else "") + "工作表。")
    return "\n".join(摘要)


# ================================================================ 主視窗（共用框架）
# 左邊：客戶切換＋選單；右邊：頁面。
#   共用的頁：「檔案與年度」（薪資檔、樣板、年度）和「人員名冊」
#   中間的頁由客戶模組自己做（例如七賢的「獎金對應」、澤利的「職稱 → 會計科目」）
BG = "#FBF6EE"
CARD_BG = "#FFFDF9"
INPUT_BG = "#FFF9F2"
PRIMARY = "#D7B89F"
PRIMARY_DARK = "#C49A7C"
SECOND_BG = "#F3E8DA"
BORDER = "#EADCCB"
TEXT = "#5A4636"
TEXT_GRAY = "#A69486"
OK_COLOR = "#6E8F5E"
WARN_COLOR = "#C08A3E"
ERR_COLOR = "#C0625A"
基準寬, 基準高 = 1000, 780
字級 = {"標題": (18, "bold"), "區塊": (12, "bold"), "標籤": (11, "bold"), "內文": (11, "normal"),
        "輸入": (11, "normal"), "小": (9, "normal"), "按鈕": (12, "bold")}


def 主視窗():
    import tkinter as tk
    from tkinter import ttk, filedialog, messagebox, simpledialog
    import tkinter.font as tkfont
    from types import SimpleNamespace

    root = tk.Tk()
    root.title("薪資印領清冊")
    root.geometry(f"{基準寬}x{基準高}")
    root.minsize(840, 620)
    root.configure(bg=BG)
    字族 = next((f for f in ("Microsoft JhengHei", "Microsoft JhengHei UI", "PingFang TC", "Heiti TC",
                            "Noto Sans CJK TC") if f in tkfont.families()), "TkDefaultFont")
    字 = {k: tkfont.Font(root, family=字族, size=s, weight=w) for k, (s, w) in 字級.items()}
    root.option_add("*TCombobox*Listbox.font", 字["輸入"])
    root.option_add("*TCombobox*Listbox.background", CARD_BG)
    root.option_add("*TCombobox*Listbox.foreground", TEXT)
    root.option_add("*TCombobox*Listbox.selectBackground", PRIMARY)
    root.option_add("*TCombobox*Listbox.selectForeground", "white")
    樣式 = ttk.Style(root)
    try:
        樣式.theme_use("clam")
    except tk.TclError:
        pass
    樣式.configure("奶茶.TCombobox", fieldbackground=INPUT_BG, background=CARD_BG, foreground=TEXT,
                 bordercolor=BORDER, lightcolor=INPUT_BG, darkcolor=INPUT_BG, arrowcolor=TEXT, arrowsize=14, padding=5)
    樣式.map("奶茶.TCombobox", fieldbackground=[("readonly", INPUT_BG), ("focus", "#FFFFFF")],
           bordercolor=[("focus", PRIMARY)], lightcolor=[("focus", "#FFFFFF")])
    樣式.configure("奶茶.Vertical.TScrollbar", background=SECOND_BG, troughcolor=CARD_BG,
                 bordercolor=CARD_BG, arrowcolor=TEXT_GRAY, lightcolor=SECOND_BG, darkcolor=SECOND_BG)
    樣式.configure("奶茶.TCheckbutton", background=CARD_BG, foreground=TEXT, font=字["內文"])
    樣式.map("奶茶.TCheckbutton", background=[("active", CARD_BG)])

    狀態 = {"設": None, "外掛": {}, "比例": 1.0, "排程": None}
    資料 = {"wb": None, "wbf": None, "工作表": [], "樣板欄": [], "樣板左欄": [], "樣板": None}
    會換行 = []

    # ---------- 小零件（客戶模組也用這些） ----------
    def 區塊標題(parent, 文字, 右邊=None, bg=BG):
        框 = tk.Frame(parent, bg=bg)
        框.pack(fill="x", padx=28, pady=(14, 6))
        tk.Frame(框, bg=PRIMARY, width=5).pack(side="left", fill="y", padx=(0, 8))
        tk.Label(框, text=文字, font=字["區塊"], fg=TEXT, bg=bg).pack(side="left")
        if 右邊:
            右邊(框)
        return 框

    def 卡片(parent, 撐滿=False):
        外框 = tk.Frame(parent, bg=BORDER)
        外框.pack(fill="both" if 撐滿 else "x", expand=撐滿, padx=28)
        內 = tk.Frame(外框, bg=CARD_BG, padx=20, pady=12)
        內.pack(fill="both", expand=True, padx=1, pady=1)
        return 內

    def 標籤(parent, 文字, 粗=True):
        return tk.Label(parent, text=文字, font=字["標籤" if 粗 else "內文"], fg=TEXT, bg=CARD_BG)

    def 小字(parent, 文字="", 色=TEXT_GRAY, 寬=560):
        l = tk.Label(parent, text=文字, font=字["小"], fg=色, bg=CARD_BG, anchor="w", justify="left")
        會換行.append((l, 寬))
        l.configure(wraplength=int(寬 * 狀態["比例"]))
        return l

    def 輸入框(parent, 變數, 寬=20):
        外框 = tk.Frame(parent, bg=BORDER)
        e = tk.Entry(外框, textvariable=變數, font=字["輸入"], bg=INPUT_BG, fg=TEXT, insertbackground=TEXT,
                     relief="flat", bd=0, width=寬)
        e.pack(fill="x", padx=1, pady=1, ipady=5, ipadx=6)
        e.bind("<FocusIn>", lambda ev: ev.widget.config(bg="#FFFFFF"), add="+")
        e.bind("<FocusOut>", lambda ev: ev.widget.config(bg=INPUT_BG), add="+")
        return 外框, e

    def 按鈕(parent, 文字, 指令, 主要=False):
        底, 深 = (PRIMARY, PRIMARY_DARK) if 主要 else (SECOND_BG, BORDER)
        b = tk.Button(parent, text=文字, command=指令, font=字["按鈕" if 主要 else "內文"], bg=底,
                      fg="white" if 主要 else TEXT, activebackground=深, activeforeground="white" if 主要 else TEXT,
                      relief="flat", bd=0, cursor="hand2", padx=14, pady=3)
        b.bind("<Enter>", lambda e: b.config(bg=深))
        b.bind("<Leave>", lambda e: b.config(bg=底))
        return b

    def 文字按鈕(parent, 文字, 指令, 色=TEXT_GRAY, 移上色=TEXT, bg=CARD_BG):
        b = tk.Label(parent, text=文字, font=字["內文"], fg=色, bg=bg, cursor="hand2")
        b.bind("<Button-1>", lambda e: 指令())
        b.bind("<Enter>", lambda e: b.config(fg=移上色))
        b.bind("<Leave>", lambda e: b.config(fg=色))
        return b

    def 下拉(parent, 值們=(), 唯讀=False, 寬=None):
        cb = ttk.Combobox(parent, values=list(值們), font=字["輸入"], style="奶茶.TCombobox",
                          state="readonly" if 唯讀 else "normal", **({"width": 寬} if 寬 else {}))
        for 事件 in ("<MouseWheel>", "<Button-4>", "<Button-5>"):   # 滾輪不換選項，只捲動畫面
            cb.bind(事件, lambda e: (滾輪(e), "break")[1])
        return cb

    def 顯示檔名(e, 變數):                    # 路徑太長時顯示最後面（檔名）
        def f(*_):
            e.after_idle(lambda: e.xview_moveto(1.0))
        變數.trace_add("write", f)
        e.bind("<FocusOut>", f, add="+")
        f()

    def 選檔案(變數, 標題, 之後, 類型=None):
        起點 = os.path.dirname(變數.get()) if 變數.get() else ""
        路徑 = filedialog.askopenfilename(parent=root, title=標題, initialdir=起點,
                                        filetypes=類型 or [("Excel 檔案", "*.xlsx *.xlsm"), ("所有檔案", "*.*")])
        if 路徑:
            變數.set(路徑)
            之後()

    # ---------- 版面 ----------
    側欄 = tk.Frame(root, bg=SECOND_BG, width=210)
    側欄.pack(side="left", fill="y")
    側欄.pack_propagate(False)
    主區 = tk.Frame(root, bg=BG)
    主區.pack(side="left", fill="both", expand=True)
    底部 = tk.Frame(主區, bg=BG)
    底部.pack(side="bottom", fill="x", padx=28, pady=(10, 18))
    tk.Label(側欄, text="薪資印領清冊", font=字["區塊"], fg=TEXT, bg=SECOND_BG).pack(anchor="w", padx=18, pady=(20, 10))
    tk.Label(側欄, text="客戶", font=字["小"], fg=TEXT_GRAY, bg=SECOND_BG).pack(anchor="w", padx=18)
    客選 = None                                      # 下面建立（下拉要用到 滾輪）
    選單框 = tk.Frame(側欄, bg=SECOND_BG)
    頁們 = OrderedDict()                             # 鍵 → {外, 內, 畫, 標題, 核心}
    選單 = {}
    目前頁 = {"名": None}

    def 滾輪(e):
        p = 頁們.get(目前頁["名"])
        if not p:
            return
        畫 = p["畫"]
        if 畫.bbox("all") is None or 畫.bbox("all")[3] <= 畫.winfo_height():
            return
        if e.num == 4 or getattr(e, "delta", 0) > 0:
            畫.yview_scroll(-2, "units")
        elif e.num == 5 or getattr(e, "delta", 0) < 0:
            畫.yview_scroll(2, "units")
    for 事件 in ("<MouseWheel>", "<Button-4>", "<Button-5>"):
        root.bind_all(事件, 滾輪, add="+")

    客選 = 下拉(側欄, 唯讀=True)
    客選.pack(fill="x", padx=18, pady=(2, 6))
    客鈕 = tk.Frame(側欄, bg=SECOND_BG)
    客鈕.pack(fill="x", padx=18, pady=(0, 6))
    客鈕.columnconfigure(0, weight=1)
    客鈕.columnconfigure(1, weight=1)
    模組說明 = tk.Label(側欄, font=字["小"], fg=TEXT_GRAY, bg=SECOND_BG, anchor="w", justify="left", wraplength=180)
    模組說明.pack(fill="x", padx=18, pady=(0, 12))
    選單框.pack(fill="x")

    def 建頁(鍵, 標題, 核心=False):
        外 = tk.Frame(主區, bg=BG)
        畫 = tk.Canvas(外, bg=BG, highlightthickness=0, bd=0)
        軸 = ttk.Scrollbar(外, orient="vertical", command=畫.yview, style="奶茶.Vertical.TScrollbar")
        畫.configure(yscrollcommand=軸.set)
        軸.pack(side="right", fill="y")
        畫.pack(side="left", fill="both", expand=True)
        內 = tk.Frame(畫, bg=BG)
        窗 = 畫.create_window((0, 0), window=內, anchor="nw")
        內.bind("<Configure>", lambda e: 畫.configure(scrollregion=畫.bbox("all")))
        畫.bind("<Configure>", lambda e: 畫.itemconfigure(窗, width=e.width))
        頁們[鍵] = {"外": 外, "內": 內, "畫": 畫, "標題": 標題, "核心": 核心}
        return 內

    def 顯示頁(名):
        if 名 not in 頁們:
            return
        if 目前頁["名"] in 頁們:
            頁們[目前頁["名"]]["外"].pack_forget()
        頁們[名]["外"].pack(fill="both", expand=True)
        目前頁["名"] = 名
        for k, m in 選單.items():
            選 = k == 名
            for w in (m["框"], m["字"], m["狀"]):
                w.configure(bg=CARD_BG if 選 else SECOND_BG)
            m["條"].configure(bg=PRIMARY if 選 else SECOND_BG)
            m["字"].configure(font=字["標籤"] if 選 else 字["內文"])

    def 畫選單():
        for w in 選單框.winfo_children():
            w.destroy()
        選單.clear()
        順序 = ["檔案"] + [k for k in 頁們 if not 頁們[k]["核心"]] + ["名冊"]
        for k in 順序:
            外框 = tk.Frame(選單框, bg=SECOND_BG, cursor="hand2")
            外框.pack(fill="x")
            條 = tk.Frame(外框, bg=SECOND_BG, width=5)
            條.pack(side="left", fill="y")
            框 = tk.Frame(外框, bg=SECOND_BG, padx=14, pady=8)
            框.pack(side="left", fill="x", expand=True)
            字標 = tk.Label(框, text=頁們[k]["標題"], font=字["內文"], fg=TEXT, bg=SECOND_BG, anchor="w")
            字標.pack(fill="x")
            狀 = tk.Label(框, text="", font=字["小"], fg=TEXT_GRAY, bg=SECOND_BG, anchor="w")
            狀.pack(fill="x")
            選單[k] = {"框": 框, "條": 條, "字": 字標, "狀": 狀}
            for w in (外框, 框, 字標, 狀):
                w.bind("<Button-1>", lambda e, x=k: 顯示頁(x))
        目前頁["名"], 舊 = None, 目前頁["名"]
        for p in 頁們.values():
            p["外"].pack_forget()
        顯示頁(舊 if 舊 in 頁們 else "檔案")

    # ================= 共用頁：檔案與年度 =================
    檔頁 = 建頁("檔案", "檔案與年度", 核心=True)
    區塊標題(檔頁, "檔案與年度")
    卡 = 卡片(檔頁)
    卡.columnconfigure(1, weight=1)
    薪資檔變數, 樣板變數, 年變數 = tk.StringVar(), tk.StringVar(), tk.StringVar()
    標籤(卡, "薪資檔").grid(row=0, column=0, sticky="w", padx=(0, 14), pady=4)
    框, e = 輸入框(卡, 薪資檔變數)
    框.grid(row=0, column=1, sticky="ew", pady=4)
    顯示檔名(e, 薪資檔變數)
    按鈕(卡, "選擇…", lambda: 選檔案(薪資檔變數, "選擇這個客戶的薪資檔", 載入薪資檔,
                                                       [("Excel 檔案", "*.xlsx *.xlsm *.xls"), ("所有檔案", "*.*")])).grid(row=0, column=2, sticky="ew", padx=(8, 0))
    標籤(卡, "清冊樣板").grid(row=1, column=0, sticky="w", padx=(0, 14), pady=4)
    框, e = 輸入框(卡, 樣板變數)
    框.grid(row=1, column=1, sticky="ew", pady=4)
    顯示檔名(e, 樣板變數)
    按鈕(卡, "選擇…", lambda: 選檔案(樣板變數, "選擇清冊樣板", 讀樣板)).grid(row=1, column=2, sticky="ew", padx=(8, 0))
    樣提示 = 小字(卡)
    樣提示.grid(row=2, column=1, sticky="w")
    年列 = tk.Frame(卡, bg=CARD_BG)
    年列.grid(row=3, column=0, columnspan=3, sticky="w", pady=(6, 0))
    標籤(年列, "清冊年度").pack(side="left", padx=(0, 10))
    tk.Spinbox(年列, from_=100, to=200, width=5, textvariable=年變數, font=字["輸入"], bg=INPUT_BG, fg=TEXT,
               insertbackground=TEXT, relief="flat", bd=0, buttonbackground=SECOND_BG, highlightthickness=1,
               highlightbackground=BORDER, highlightcolor=PRIMARY, justify="center").pack(side="left", ipady=4)
    標籤(年列, "年", 粗=False).pack(side="left", padx=(6, 0))
    外掛區 = tk.Frame(檔頁, bg=BG)                  # 客戶模組可以在檔案頁下面加自己的卡片（例如前期薪資）
    外掛區.pack(fill="x")

    def 目前年():
        try:
            return int(年變數.get())
        except ValueError:
            return None

    # ================= 共用頁：人員名冊 =================
    名頁 = 建頁("名冊", "人員名冊", 核心=True)
    名冊變數 = tk.StringVar()
    區塊標題(名頁, "人員名冊（選填）")
    名卡 = 卡片(名頁)
    名卡.columnconfigure(1, weight=1)
    標籤(名卡, "人員名冊").grid(row=0, column=0, sticky="w", padx=(0, 14), pady=4)
    框, e = 輸入框(名卡, 名冊變數)
    框.grid(row=0, column=1, sticky="ew", pady=4)
    顯示檔名(e, 名冊變數)
    按鈕(名卡, "選擇…", lambda: 選檔案(名冊變數, "選擇人員名冊（清冊的基本資料以這份為準）", 更新名冊狀態,
                                    [("Excel 檔案", "*.xlsx *.xlsm *.xls"), ("所有檔案", "*.*")])).grid(
        row=0, column=2, sticky="ew", padx=(8, 0))
    文字按鈕(名卡, "✕", lambda: (名冊變數.set(""), 更新名冊狀態()), 移上色=ERR_COLOR).grid(row=0, column=3, padx=(8, 0))
    名冊提示 = 小字(名卡, 寬=520)
    名冊提示.grid(row=1, column=1, columnspan=3, sticky="w")
    e.bind("<Return>", lambda ev: 更新名冊狀態())
    區塊標題(名頁, "選了人員名冊會怎樣")
    說卡 = 卡片(名頁)
    名說明 = tk.Label(說卡, font=字["內文"], fg=TEXT, bg=CARD_BG, anchor="w", justify="left", text=(
        "1. 清冊上的編號、身分證號、地址，改用人員名冊上的資料（客戶給的不採用）\n"
        "   客戶模組有自己的編號規則的話（例如依科目編號），編號照那個規則\n"
        "2. 先用姓名找名冊上的人，找不到再用客戶的身分證找；名冊上沒有的人才用客戶的資料\n"
        "3. 輸出檔多一張「人員名冊核對」：身分證不同、所得類別不同、地址不同、名冊找不到…\n"
        "4. 名冊可以是 .xlsx 或舊版 .xls（舊版要先 pip install xlrd），欄位：編號、姓名、身分證、類別代號、類別、地址"))
    名說明.pack(fill="x")
    會換行.append((名說明, 660))

    def 更新名冊狀態():
        p = 名冊變數.get().strip()
        if not p:
            名冊提示.configure(text="選填｜選了的話，清冊的編號、身分證、地址以這份為準，並跟客戶資料核對", fg=TEXT_GRAY)
        elif not os.path.exists(p):
            名冊提示.configure(text="✗ 找不到這個檔案", fg=ERR_COLOR)
        else:
            try:
                人 = 讀前年度(p)
                有證 = sum(1 for v in 人.values() if v.get("身分證號"))
                名冊提示.configure(text=f"✓ 讀到 {len(人)} 人（{有證} 人有身分證號）｜清冊的編號、身分證、地址會用這份", fg=OK_COLOR)
            except SystemExit as ex:
                名冊提示.configure(text=f"✗ {ex.code}", fg=ERR_COLOR)
            except Exception as ex:
                名冊提示.configure(text=f"✗ 讀不到：{ex}", fg=ERR_COLOR)

    # ---------- 讀檔 ----------
    def 叫外掛(名, *參數):
        f = 狀態["外掛"].get(名)
        if f:
            try:
                return f(*參數)
            except tk.TclError:
                pass

    def 讀樣板():
        路徑 = 樣板變數.get()
        資料["樣板欄"], 資料["樣板左欄"], 資料["樣板"] = [], [], None
        if not 路徑:
            樣提示.configure(text="", fg=TEXT_GRAY)
        elif not os.path.exists(路徑):
            樣提示.configure(text="✗ 找不到檔案", fg=ERR_COLOR)
        else:
            try:
                ws = 樣板工作表(openpyxl.load_workbook(路徑))
                樣 = 樣板(ws, 要稅=bool(拆詞((狀態["設"] or {}).get("扣繳稅額", ""))))
                lab = lambda i: 清字(ws.cell(樣.首列 + i, 樣.明細欄).value) or "（自動）"
                資料["樣板"] = 樣
                資料["樣板欄"] = [k for k, c in 樣.欄.items() if c > 樣.明細欄 and k not in 中文月 and k != "合計"]
                資料["樣板左欄"] = [k for k, c in 樣.欄.items() if c < 樣.明細欄]
                列名 = "、".join(lab(i) for i in range(樣.組長)) if 樣.舊末列 != 樣.表頭列 else \
                    "樣板是空白的，會自動做：薪資總額、退休自提" + ("、扣繳稅額" if "稅" in 樣.角色 else "")
                樣提示.configure(text=f"✓ 每人 {樣.組長} 列（{列名}）｜月份欄 {sum(1 for k in 樣.欄 if k in 中文月)} 個",
                              fg=OK_COLOR)
            except SystemExit as e:
                樣提示.configure(text=f"✗ {e.code}", fg=ERR_COLOR)
            except Exception as e:
                樣提示.configure(text=f"✗ 讀不到：{e}", fg=ERR_COLOR)
        叫外掛("樣板")

    def 載入薪資檔(重讀=True):
        路徑 = 薪資檔變數.get()
        if 重讀:
            資料.update(wb=None, wbf=None, 工作表=[])
            if 路徑 and os.path.exists(路徑) and not os.path.isdir(路徑):
                try:
                    root.configure(cursor="watch"); root.update()
                    if 路徑.lower().endswith(".xls"):
                        資料["wb"] = 資料["wbf"] = 開活頁簿(路徑)
                    else:
                        資料["wb"] = openpyxl.load_workbook(路徑, data_only=True)
                        資料["wbf"] = openpyxl.load_workbook(路徑)
                    資料["工作表"] = [n.strip() for n in 資料["wb"].sheetnames]
                except Exception as e:
                    messagebox.showerror("讀取失敗", f"薪資檔讀不到：{e}", parent=root)
                finally:
                    root.configure(cursor="")
        叫外掛("載入")

    def 年改了(*_):
        叫外掛("年改了")
        載入薪資檔(重讀=False)
    年變數.trace_add("write", 年改了)

    # ---------- 設定 ⇄ 畫面 ----------
    def 收集():
        設 = copy.deepcopy(狀態["設"] or 空白設定)
        設.update({"年度": 目前年() or 設.get("年度", 115), "薪資檔": 薪資檔變數.get(), "樣板": 樣板變數.get(),
                  "人員名冊": 名冊變數.get().strip()})
        叫外掛("收集", 設)
        return 設

    介 = SimpleNamespace(tk=tk, ttk=ttk, root=root, 字=字, 資料=資料, 會換行=會換行, 區塊標題=區塊標題, 卡片=卡片,
                       標籤=標籤, 小字=小字, 輸入框=輸入框, 按鈕=按鈕, 文字按鈕=文字按鈕, 下拉=下拉, 顯示檔名=顯示檔名,
                       選檔案=選檔案, 目前年=目前年, 收集=收集, 檔案頁=外掛區, 加頁=lambda 鍵, 標題: 建頁(鍵, 標題),
                       薪資檔變數=薪資檔變數, 重讀=lambda: 載入薪資檔(重讀=False), messagebox=messagebox,
                       BG=BG, CARD_BG=CARD_BG, INPUT_BG=INPUT_BG, PRIMARY=PRIMARY, PRIMARY_DARK=PRIMARY_DARK,
                       SECOND_BG=SECOND_BG, BORDER=BORDER, TEXT=TEXT, TEXT_GRAY=TEXT_GRAY, OK=OK_COLOR,
                       WARN=WARN_COLOR, ERR=ERR_COLOR)

    def 放進畫面(設):
        狀態["設"] = 設
        # 換掉客戶模組的頁
        for k in [k for k in 頁們 if not 頁們[k]["核心"]]:
            頁們[k]["外"].destroy()
            del 頁們[k]
        for w in 外掛區.winfo_children():
            w.destroy()
        會換行[:] = [(w, x) for w, x in 會換行 if w.winfo_exists()]
        m = 取模組(設.get("模組", "通用"))
        模組說明.configure(text=f"讀法：{getattr(m, '名稱', 設.get('模組', ''))}" if m else "✗ 找不到客戶模組")
        年變數.set(str(設.get("年度", 115)))
        薪資檔變數.set(設.get("薪資檔", "")); 樣板變數.set(設.get("樣板", "")); 名冊變數.set(設.get("人員名冊", ""))
        狀態["外掛"] = {}
        try:
            狀態["外掛"] = (m.頁面(介, 設) or {}) if m is not None and hasattr(m, "頁面") else {}
        except Exception as e:              # 客戶模組的畫面有錯：顯示錯誤，其他頁照常
            import traceback
            traceback.print_exc()
            messagebox.showerror("客戶模組有錯", f"{設.get('模組')}.py 的畫面出錯：\n{type(e).__name__}: {e}", parent=root)
        畫選單()
        讀樣板()
        載入薪資檔()
        更新名冊狀態()
        for p in 頁們.values():
            p["畫"].yview_moveto(0)
        顯示頁("檔案")

    # ---------- 客戶：切換、新增、複製、改名、刪除 ----------
    def 畫客戶(選=None):
        名們 = 客戶清單()
        客選.configure(values=名們)
        選 = 選 if 選 in 名們 else (名們[0] if 名們 else "")
        客選.set(選)
        return 選

    def 換客戶(e=None):
        名 = 客選.get()
        if 狀態["設"] and 狀態["設"]["客戶"] != 名:
            存客戶(收集())
        if 名:
            上次客戶(名)
            放進畫面(讀客戶(名))
    客選.bind("<<ComboboxSelected>>", 換客戶)

    def 問名稱(標題, 預設=""):
        名 = simpledialog.askstring(標題, "客戶名稱：", initialvalue=預設, parent=root)
        if not 名 or not 名.strip():
            return None
        名 = 名.strip()
        if 名 in 客戶清單():
            messagebox.showerror("名稱重複", f"已經有「{名}」這個客戶了", parent=root)
            return None
        return 名

    def 選模組(預設="通用"):                      # 新增客戶時選「讀法」（哪個客戶模組）
        模組們 = 載入模組們()
        if len(模組們) <= 1:
            return next(iter(模組們), "通用")
        窗 = tk.Toplevel(root, bg=CARD_BG, padx=20, pady=16)
        窗.title("選擇讀法")
        窗.transient(root)
        tk.Label(窗, text="這個客戶的薪資檔要用哪一種讀法？", font=字["標籤"], fg=TEXT, bg=CARD_BG).pack(anchor="w", pady=(0, 8))
        選 = tk.StringVar(value=預設 if 預設 in 模組們 else next(iter(模組們)))
        for 鍵, m in 模組們.items():
            tk.Radiobutton(窗, text=f"{getattr(m, '名稱', 鍵)}　—　{getattr(m, '說明', '')}", variable=選, value=鍵,
                           font=字["內文"], fg=TEXT, bg=CARD_BG, activebackground=CARD_BG, selectcolor=INPUT_BG,
                           anchor="w", justify="left", wraplength=520).pack(fill="x", anchor="w")
        結果 = {"鍵": None}
        列 = tk.Frame(窗, bg=CARD_BG)
        列.pack(fill="x", pady=(12, 0))
        按鈕(列, "確定", lambda: (結果.update(鍵=選.get()), 窗.destroy()), 主要=True).pack(side="right")
        按鈕(列, "取消", 窗.destroy).pack(side="right", padx=(0, 8))
        窗.grab_set()
        root.wait_window(窗)
        return 結果["鍵"]

    def 新增客戶():
        名 = 問名稱("新增客戶")
        if not 名:
            return
        鍵 = 選模組()
        if not 鍵:
            return
        if 狀態["設"]:
            存客戶(收集())
        設 = 新客戶設定(名, 鍵)
        設.update(年度=目前年() or 115, 樣板=樣板變數.get())
        存客戶(設)
        畫客戶(名)
        放進畫面(設)

    def 複製客戶():
        if not 狀態["設"]:
            return
        名 = 問名稱("複製客戶", 狀態["設"]["客戶"] + "（複製）")
        if not 名:
            return
        存客戶(收集())
        設 = 收集()
        設.update(客戶=名, 薪資檔="")
        存客戶(設)
        畫客戶(名)
        放進畫面(設)

    def 改名客戶():
        if not 狀態["設"]:
            return
        舊 = 狀態["設"]["客戶"]
        名 = 問名稱("改名", 舊)
        if not 名:
            return
        設 = 收集()
        設["客戶"] = 名
        存客戶(設)
        刪客戶(舊)
        狀態["設"] = 設
        畫客戶(名)
        上次客戶(名)

    def 刪除客戶():
        if not 狀態["設"]:
            return
        名 = 狀態["設"]["客戶"]
        if not messagebox.askyesno("刪除客戶", f"確定要刪除「{名}」的設定嗎？（不會動到任何 Excel 檔）", parent=root):
            return
        刪客戶(名)
        狀態["設"] = None
        選 = 畫客戶()
        if 選:
            放進畫面(讀客戶(選))

    for i, (文, 指令) in enumerate((("新增", 新增客戶), ("複製", 複製客戶), ("改名", 改名客戶), ("刪除", 刪除客戶))):
        按鈕(客鈕, 文, 指令).grid(row=i // 2, column=i % 2, sticky="ew", padx=(0 if i % 2 == 0 else 3, 3 if i % 2 == 0 else 0), pady=2)

    # ---------- 產生 ----------
    def 開始():
        try:
            設 = 收集()
            確認 = 叫外掛("產生前確認", 設)
            if 確認 is False:
                return
            上次 = 設.get("輸出資料夾", "")
            資料夾 = filedialog.askdirectory(parent=root, title="選擇輸出資料夾",
                                          initialdir=上次 if 上次 and os.path.isdir(上次) else os.path.dirname(設["薪資檔"]))
            if not 資料夾:
                return
            設["輸出資料夾"] = 資料夾
            存客戶(設); 狀態["設"] = 設
            root.configure(cursor="watch"); root.update()
            結果 = 執行(設, 資料夾)
            messagebox.showinfo("完成", 結果, parent=root)
        except SystemExit as e:
            if e.code not in (None, 0):
                messagebox.showerror("錯誤", str(e.code), parent=root)
        except Exception as e:
            import traceback
            traceback.print_exc()
            messagebox.showerror("發生錯誤", f"{type(e).__name__}: {e}", parent=root)
        finally:
            root.configure(cursor="")

    def 只存設定():
        if 狀態["設"]:
            設 = 收集(); 存客戶(設); 狀態["設"] = 設
            messagebox.showinfo("已儲存", f"「{設['客戶']}」的設定已儲存。", parent=root)

    按鈕(底部, "儲存設定", 只存設定).pack(side="right", padx=(12, 0), ipady=5)
    按鈕(底部, "產　生", 開始, 主要=True).pack(side="left", fill="x", expand=True, ipady=7)

    # ---------- 左邊選單的狀態小字（每 0.6 秒更新） ----------
    def 更新側欄():
        try:
            年 = 目前年()
            有檔 = bool(薪資檔變數.get()) and os.path.exists(薪資檔變數.get()) and bool(樣板變數.get()) and os.path.exists(樣板變數.get())
            if "檔案" in 選單:
                選單["檔案"]["狀"].configure(text=("✓ " if 有檔 else "✗ 還沒選檔案｜") + (f"{年} 年" if 年 else "年度？"),
                                           fg=OK_COLOR if 有檔 else ERR_COLOR)
            if "名冊" in 選單:
                if not 名冊變數.get().strip():
                    選單["名冊"]["狀"].configure(text="選填｜沒有選", fg=TEXT_GRAY)
                else:
                    好 = 名冊提示.cget("text").startswith("✓")
                    選單["名冊"]["狀"].configure(text="✓ 會用名冊的基本資料" if 好 else "✗ 名冊讀不到", fg=OK_COLOR if 好 else ERR_COLOR)
            for k, (文, 色) in (叫外掛("狀態") or {}).items():
                if k in 選單:
                    選單[k]["狀"].configure(text=文, fg=色)
        except tk.TclError:
            return
        root.after(600, 更新側欄)

    # ---------- 等比例縮放 ----------
    def 套用縮放():
        狀態["排程"] = None
        比例 = max(0.85, min(2.0, min(root.winfo_width() / 基準寬, root.winfo_height() / 基準高)))
        if abs(比例 - 狀態["比例"]) < 0.03:
            return
        狀態["比例"] = 比例
        for k, (s, _) in 字級.items():
            字[k].configure(size=max(8, round(s * 比例)))
        for 元件, 寬 in 會換行:
            try:
                元件.configure(wraplength=int(寬 * 比例))
            except tk.TclError:
                pass
        樣式.configure("奶茶.TCombobox", arrowsize=int(14 * 比例), padding=int(5 * 比例))
        側欄.configure(width=int(210 * 比例))
        模組說明.configure(wraplength=int(180 * 比例))
    root.bind("<Configure>", lambda e: (e.widget is root and 狀態["排程"] is None
                                        and 狀態.update(排程=root.after(80, 套用縮放))))

    def 關閉():
        if 狀態["設"]:
            存客戶(收集())
        root.destroy()
    root.protocol("WM_DELETE_WINDOW", 關閉)

    # ---------- 開啟 ----------
    for 元件, 寬 in 會換行:
        元件.configure(wraplength=寬)
    選 = 畫客戶(上次客戶())
    if 選:
        放進畫面(讀客戶(選))
    else:
        畫選單()
    更新側欄()
    root.主視窗資料 = {"客選": 客選, "換客戶": 換客戶, "收集": 收集, "開始": 開始, "顯示頁": 顯示頁, "狀態": 狀態,
                     "薪資檔變數": 薪資檔變數, "樣板變數": 樣板變數, "名冊變數": 名冊變數, "載入薪資檔": 載入薪資檔,
                     "讀樣板": 讀樣板, "新增客戶": 新增客戶, "放進畫面": 放進畫面, "頁們": 頁們}
    return root


def main():
    if len(sys.argv) > 1:      # 指令：python 清冊共用.py 客戶名稱 [輸出資料夾]
        名 = sys.argv[1]
        if 名 not in 客戶清單():
            sys.exit(f"找不到客戶「{名}」。目前有：{'、'.join(客戶清單())}")
        設 = 讀客戶(名)
        資料夾 = sys.argv[2] if len(sys.argv) > 2 else (設.get("輸出資料夾") or os.path.dirname(設["薪資檔"]))
        print(執行(設, 資料夾))
        return
    import importlib.util
    if importlib.util.find_spec("tkinter") is None:
        sys.exit("這台電腦的 Python 沒有 tkinter，無法打開視窗。\n請改用：python 清冊共用.py 客戶名稱 輸出資料夾")
    主視窗().mainloop()


if __name__ == "__main__":
    sys.modules.setdefault("清冊共用", sys.modules[__name__])
    main()
